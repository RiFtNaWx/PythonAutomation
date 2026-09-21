"""RS622-shaped campaign sheet_map. One place for outline + known paste cells.

Do not invent photo cells. Discover merged 4-col DUT boxes or 8-col CHA/CHB pairs.
Do not write FILL_ME. Tracking-probed numbers only.
"""
from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path
from typing import Any

import yaml

from ate.core.paths import TEST_DB_ROOT
from ate.reporting.session_values import _fold_name

_FILL = frozenset({"FILL_ME", "STUB", "TODO", "TBD"})
_INTRO_KEYS = {
    "test conditions": "conditions",
    "test circuitry": "circuitry",
    "datasheet": "datasheet",
    "test conclusion": "conclusion",
}
_SKIP_SHEETS = frozenset({"Summary", "Checklist", "Mapping", "Status", "Problems"})
_NAMING = {
    "screenshot": "{TEST}_{DUT}_{VARIANT}_{TIMESTAMP}.jpg",
    "graph": "{TEST}_{DUT}_{VARIANT}_{TIMESTAMP}.png",
}
_FAM_MODE = {
    "opamp": "BUFFER",
    "logic": "LOGIC",
    "level": "LOGIC",
    "switch": "LIM_RS2323",
    "power": "LDO",
}
_KEY_MODE = {
    "gbw": "G11",
    "ort": "G_NEG100",
    "vos": "G1001",
    "vossweep": "G1001",
    "noise": "ATE",
    "psrr": "ATE",
    "cmrr": "ATE",
    "aol": "ATE",
    "emirr": "ATE",
    "iplus": "LIM_RS2323",
    "iq": "LDO",
}

# Iplus stub (SeeLim) is B2. Parameter DUT_1..n grids are probed, not B2.
_IPLUS_STUB = {"IPLUS_uA": "B2"}
_DUT_HDR = re.compile(r"^DUT_?(\d+)$")
_MID_TOKEN = re.compile(r"^[A-Za-z][A-Za-z0-9]*_[A-Za-z0-9]+$")
_TEST_DUT_MIDS: dict[str, tuple[str, ...]] = {
    "iplus": ("IPLUS_uA",),
    "leakage_off": ("IOZ_uA",),
    "leakage_on": ("ION_uA",),
    "input_leakage": ("IIN_uA",),
    "usb_ron": ("USB_RON_ohm",),
    "usb_ton_toff": ("USB_TON_ns", "USB_TOFF_ns"),
    "ron": ("RON_ohm",),
    "ton_toff": ("TON_ns", "TOFF_ns"),
    "vth": ("VTH_V",),
    "con_coff": ("CIN_pF", "CON_pF", "COFF_pF"),
    "tbbm": ("TBBM_ns",),
    "iso": ("ISO_dB",),
    "xtalk": ("XTALK_dB",),
    "usb_iso": ("USB_ISO_dB",),
    "usb_xtalk": ("USB_XTALK_dB",),
    "i2c_ii": ("II_uA",),
    "i2c_ron": ("RON_ohm",),
    "i2c_cioff": ("CIOFF_pF",),
    "cin": ("CIN_pF",),
    "cpd": ("CPD_pF",),
    "ioff_leakage": ("IOFF_uA",),
    "voh_load": (
        "VOH_2p0V_100uA",
        "VOH_3p3V_100uA",
        "VOH_4p5V_100uA",
        "VOH_5p0V_100uA",
        "VOH_5p5V_100uA",
        "VOH_2p0V",
        "VOH_3p3V",
        "VOH_4p5V",
        "VOH_5p0V",
        "VOH_5p5V",
    ),
    "vol_load": ("VOL_2p0V", "VOL_3p3V", "VOL_4p5V", "VOL_5p0V", "VOL_5p5V"),
    "tp": ("TPD_PHL_ns", "TPD_PLH_ns"),
    "clk_q": ("CLKQ_PHL_ns", "CLKQ_PLH_ns"),
    "pulse_width": ("PULSE_ns",),
    "serial_shift": ("Q7_HIGH_V", "Q7_LOW_V"),
    "delta_supply_current": ("DELTA_ICC_uA",),
    "input_leakage_sweep": ("II_uA",),
    "ioz": ("IOZ_uA",),
    "ten": ("TEN_ns",),
    "tdis": ("TDIS_ns",),
}
_SHEET_DUT_MIDS: dict[str, tuple[str, ...]] = {
    "iplus": ("IPLUS_uA",),
    "leakageoff": ("IOZ_uA",),
    "leakageon": ("ION_uA",),
    "inputleakage": ("IIN_uA",),
    "usbron": ("USB_RON_ohm",),
    "tontoff": ("USB_TON_ns", "USB_TOFF_ns"),
    "ton": ("TON_ns", "TOFF_ns"),
    "vth": ("VTH_V",),
    "cinconcoff": ("CIN_pF", "CON_pF", "COFF_pF"),
    "tbbm": ("TBBM_ns",),
    "iso": ("ISO_dB",),
    "xtalk": ("XTALK_dB",),
    "usbiso": ("USB_ISO_dB",),
    "usbxtalk": ("USB_XTALK_dB",),
    "ii": ("II_uA",),
    "i2cii": ("II_uA",),
    "i2cron": ("RON_ohm",),
    "cioff": ("CIOFF_pF",),
    "i2ccioff": ("CIOFF_pF",),
    "cin": ("CIN_pF",),
    "cpd": ("CPD_pF",),
    "ioffleakage": ("IOFF_uA",),
    "deltaidd": ("DELTA_ICC_uA",),
    "inputleakagesweep": ("II_uA",),
    "ioz": ("IOZ_uA",),
    "ten": ("TEN_ns",),
    "tdis": ("TDIS_ns",),
    "tw": ("PULSE_ns",),
    "q7": ("Q7_HIGH_V", "Q7_LOW_V"),
}
_GBW_TTSOP = {
    "GBW_MHz": {"CHA": ["R20", "S20", "T20", "U20"], "CHB": ["AA20", "AB20", "AC20", "AD20"]},
    "F_3dB_kHz": {"CHA": ["R19", "S19", "T19", "U19"], "CHB": ["AA19", "AB19", "AC19", "AD19"]},
    "VOUT_1k_mV": {"CHA": ["R17", "S17", "T17", "U17"], "CHB": ["AA17", "AB17", "AC17", "AD17"]},
}
_GBW_SOP8 = {
    "GBW_MHz": {"CHA": ["C21", "D21", "E21"], "CHB": ["J21", "K21", "L21"]},
    "F_3dB_kHz": {"CHA": ["C20", "D20", "E20"], "CHB": ["J20", "K20", "L20"]},
    "VOUT_1k_mV": {"CHA": ["C18", "D18", "E18"], "CHB": ["J18", "K18", "L18"]},
}
_VOS_TTSOP = {
    "VOS_mV": {"CHA": ["R16", "S16", "T16", "U16"], "CHB": ["AA16", "AB16", "AC16", "AD16"]}
}
_VOS_SOP8 = {
    "VOS_mV": {"CHA": ["B16", "C16", "D16", "E16"], "CHB": ["I16", "J16", "K16", "L16"]}
}

_GRID_COLS = ("A", "E", "I", "M", "Q", "U", "Y", "AC")
_CELL_ROW = re.compile(r"^([A-Z]{1,3})(\d+)$")


def _is_copied_dut_grid(photos: dict[str, Any]) -> bool:
    """Same-row 8/16-box A/E/I/M grid copied from one TTSOP snapshot."""
    if not isinstance(photos, dict) or len(photos) not in (8, 16):
        return False
    rows: set[str] = set()
    cols: set[str] = set()
    for raw in photos.values():
        m = _CELL_ROW.match(str(raw).strip().upper())
        if not m:
            return False
        cols.add(m.group(1))
        rows.add(m.group(2))
    return cols <= set(_GRID_COLS) and 1 <= len(rows) <= 2


def attach_known_photos(tests: dict[str, Any], sheets: list[str], *, wb: Any = None) -> None:
    """Photo cells = merged 4-col boxes on the live sheet. Do not invent A70/A45."""
    if wb is None:
        return
    from ate.reporting.golden_layout import discover_photo_anchors

    sheet_set = {str(s) for s in sheets} if sheets else set(getattr(wb, "sheetnames", []))
    for name in list(getattr(wb, "sheetnames", []) or []):
        if sheet_set and name not in sheet_set:
            continue
        entry = _pick_entry(tests, name, _folder_key(name))
        if entry is None:
            continue
        found = discover_photo_anchors(wb[name])
        paste = entry.get("paste")
        if not isinstance(paste, dict):
            paste = {}
            entry["paste"] = paste
        if found:
            paste["photos"] = found
            continue
        photos = paste.get("photos")
        if _is_copied_dut_grid(photos if isinstance(photos, dict) else {}):
            paste.pop("photos", None)
        if not paste:
            entry.pop("paste", None)


def attach_known_narrative(tests: dict[str, Any], sheets: list[str], *, wb: Any = None) -> None:
    """Intro B cells: Test Conditions / Circuitry / Datasheet / Conclusion. Probe, do not guess."""
    if wb is None:
        return
    sheet_set = {str(s) for s in sheets} if sheets else set(getattr(wb, "sheetnames", []))
    for name in list(getattr(wb, "sheetnames", []) or []):
        if sheet_set and name not in sheet_set:
            continue
        entry = _pick_entry(tests, name, _folder_key(name))
        if entry is None:
            continue
        found: dict[str, str] = {}
        ws = wb[name]
        for row in range(1, 81):
            lab = str(ws.cell(row, 1).value or "").strip().lower()
            key = _INTRO_KEYS.get(lab)
            if not key:
                continue
            found[key] = f"B{row}"
        if not found:
            continue
        paste = entry.get("paste")
        if not isinstance(paste, dict):
            paste = {}
            entry["paste"] = paste
        paste["narrative"] = found


TEST_OUTLINE = ("folder", "excel_sheet", "fixture_mode", "automated", "dut_iterations")


def _is_fill(raw: Any) -> bool:
    if isinstance(raw, str) and raw.split("#", 1)[0].strip().upper() in _FILL:
        return True
    if isinstance(raw, dict):
        return bool(raw) and all(_is_fill(v) for v in raw.values())
    if isinstance(raw, (list, tuple)):
        return bool(raw) and all(_is_fill(v) for v in raw)
    return False


def _folder_key(sheet: str) -> str:
    raw = re.sub(r"[^\w\-]+", "", (sheet or "").replace(" ", ""))
    return raw or "Sheet"


def _family_from_component(name: str) -> str:
    n = str(name or "").strip().lower().replace(" ", "")
    if n in ("opamp", "opa"):
        return "opamp"
    if n in ("analogswitch", "switch", "lim"):
        return "switch"
    if n == "power":
        return "power"
    if n == "level":
        return "level"
    if n == "logic":
        return "logic"
    return n


def _infer_mode(key: str, family: str) -> str:
    folds = _fold_name(key)
    for name, mode in _KEY_MODE.items():
        if name in folds:
            return mode
    return _FAM_MODE.get(family, "LOGIC")


def _cell_ok(ws: Any, coord: str) -> bool:
    try:
        from openpyxl.cell.cell import MergedCell

        return not isinstance(ws[coord], MergedCell)
    except Exception:
        return False


def _cell_num(ws: Any, coord: str) -> bool:
    if not _cell_ok(ws, coord):
        return False
    try:
        v = ws[coord].value
    except Exception:
        return False
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _as_float(raw: Any) -> float | None:
    if raw is None or isinstance(raw, bool):
        return None
    if isinstance(raw, (int, float)):
        return float(raw)
    try:
        return float(str(raw).strip().replace("V", "").strip())
    except (TypeError, ValueError):
        return None


def _banner_row(ws: Any, label: str) -> int | None:
    want = label.strip().upper()
    for r in range(1, 81):
        if str(ws.cell(r, 1).value or "").strip().upper() == want:
            return r
    return None


def _vcc_token(vcc: float) -> str:
    s = f"{vcc:.2f}".rstrip("0").rstrip(".")
    if "." not in s:
        s = f"{s}p0"
    else:
        s = s.replace(".", "p")
    return s


def _probe_vox_block(ws: Any, banner: str, prefix: str) -> dict[str, list[str]]:
    """DUT G/H/I for each unique Vcc under a VOH/VOL banner. First row wins on dup 3.0."""
    start = _banner_row(ws, banner)
    if not start:
        return {}
    out: dict[str, list[str]] = {}
    for r in range(start + 1, min(start + 16, 120)):
        n = _as_float(ws.cell(r, 1).value)
        if n is None:
            continue
        sid = f"{prefix}_{_vcc_token(n)}V"
        if sid in out:
            continue
        out[sid] = [f"G{r}", f"H{r}", f"I{r}"]
    return out


def _probe_vox_corners(ws: Any) -> dict[str, list[str]]:
    """Sheet Vcc rows only. Do not alias 1.65 as VOH_2p0V."""
    out = _probe_vox_block(ws, "VOH", "VOH")
    out.update(_probe_vox_block(ws, "VOL", "VOL"))
    return out


# DUT1 A1 on RS0204-style lab sheets. Skip tr/tf: A1 there is stimulus edge time.
_A1_IDS: dict[str, str] = {
    "VIH": "VIH_V",
    "VIL": "VIL_V",
    "Il": "IL_uA",
}


def _probe_dut_value_cells(ws: Any) -> list[str]:
    """Header DUT_1..n; values on the next row. Empty if the sheet has no DUT columns."""
    from openpyxl.utils import get_column_letter

    header_row = 0
    cols: list[tuple[int, int]] = []
    for r in range(1, 16):
        found: list[tuple[int, int]] = []
        for c in range(1, 24):
            raw = str(ws.cell(r, c).value or "").strip().upper().replace(" ", "")
            m = _DUT_HDR.match(raw)
            if m:
                found.append((int(m.group(1)), c))
        if found:
            header_row = r
            cols = found
            break
    if not header_row:
        return []
    vr = header_row + 1
    return [f"{get_column_letter(c)}{vr}" for _n, c in sorted(cols)]


def _shift_row(cells: list[str], delta: int) -> list[str]:
    out: list[str] = []
    for tok in cells:
        m = _CELL_ROW.match(str(tok or "").strip().upper())
        if not m:
            continue
        out.append(f"{m.group(1)}{int(m.group(2)) + delta}")
    return out


def _probe_named_dut_mids(ws: Any) -> dict[str, list[str]]:
    """A-column measurement ids on DUT rows (USB_TON_ns / USB_TOFF_ns)."""
    base = _probe_dut_value_cells(ws)
    if not base:
        return {}
    m0 = _CELL_ROW.match(base[0])
    if not m0:
        return {}
    start = int(m0.group(2))
    letters = []
    for tok in base:
        mm = _CELL_ROW.match(tok)
        if mm:
            letters.append(mm.group(1))
    found: dict[str, list[str]] = {}
    for r in range(start, start + 8):
        raw = str(ws.cell(r, 1).value or "").strip()
        if not _MID_TOKEN.match(raw):
            continue
        found[raw] = [f"{col}{r}" for col in letters]
    return found


def _mids_for_entry(key: str, entry: dict[str, Any]) -> tuple[str, ...]:
    names = (key, str(entry.get("folder") or ""), str(entry.get("excel_sheet") or ""))
    for raw in names:
        n = re.sub(r"[^a-z0-9]+", "", str(raw or "").lower())
        if n in _SHEET_DUT_MIDS:
            return _SHEET_DUT_MIDS[n]
        for tid, mids in _TEST_DUT_MIDS.items():
            if n == re.sub(r"[^a-z0-9]+", "", tid):
                return mids
    return ()


def _label_stub_mid(ws: Any, mid: str) -> bool:
    """Write A2 on a SeeLim-style '(stub)' sheet. Do not touch Parameter DUT grids."""
    a1 = str(ws["A1"].value or "").strip().lower()
    if a1 == "parameter" or _probe_dut_value_cells(ws):
        return False
    if not ws["A2"].value:
        ws["A2"] = f"ATE {mid}"
        return True
    return False


def _label_iplus_stub(ws: Any) -> bool:
    """Write A2 label only on the B2 stub sheet. Do not touch Parameter DUT grids."""
    return _label_stub_mid(ws, "IPLUS_uA")


def _probe_a1_cell(ws: Any) -> str:
    """Header 'A1' then the cell immediately below (DUT1 value). Probed, not guessed."""
    from openpyxl.utils import get_column_letter

    for r in range(1, 50):
        for c in range(1, 24):
            raw = str(ws.cell(r, c).value or "").strip().upper().replace(" ", "")
            if raw == "A1":
                return f"{get_column_letter(c)}{r + 1}"
    return ""


def _probe_max_ua(ws: Any, sid: str) -> dict[str, str]:
    """Cell to the right of a short 'maximum (uA)' banner. Skip conclusion prose."""
    from openpyxl.utils import get_column_letter

    for r in range(1, 40):
        for c in range(1, 8):
            t = str(ws.cell(r, c).value or "").lower().replace("μ", "u").replace("µ", "u")
            if len(t) < 40 and re.search(r"maximum\s*\(\s*ua\s*\)", t):
                return {sid: f"{get_column_letter(c + 1)}{r}"}
    return {}


def _probe_icc_max(ws: Any) -> dict[str, str]:
    return _probe_max_ua(ws, "ICC_uA")


def _probe_col_first_value(ws: Any, *needles: str) -> str:
    """Header contains all needles; first numeric cell below that column."""
    from openpyxl.utils import get_column_letter

    want = tuple(str(n or "").lower() for n in needles if n)
    if not want:
        return ""
    for r in range(1, 40):
        for c in range(1, 16):
            t = str(ws.cell(r, c).value or "").lower()
            if len(t) >= 48 or not all(n in t for n in want):
                continue
            right = ws.cell(r, c + 1).value
            if isinstance(right, str) and right.strip().startswith("="):
                return f"{get_column_letter(c + 1)}{r}"
            for rr in range(r + 1, min(r + 16, 80)):
                v = ws.cell(rr, c).value
                if v is None:
                    continue
                if isinstance(v, (int, float)):
                    return f"{get_column_letter(c)}{rr}"
                s = str(v).strip()
                if not s:
                    continue
                if s.startswith("="):
                    return f"{get_column_letter(c)}{rr}"
                try:
                    float(s.replace(",", ""))
                except ValueError:
                    continue
                return f"{get_column_letter(c)}{rr}"
            if isinstance(right, (int, float)):
                return f"{get_column_letter(c + 1)}{r}"
            rs = str(right or "").strip()
            if rs:
                try:
                    float(rs.replace(",", ""))
                    return f"{get_column_letter(c + 1)}{r}"
                except ValueError:
                    pass
            continue
    return ""


def _path_kind(xlsx: Path | None) -> str:
    s = str(xlsx or "").replace("/", "\\").upper()
    if "TTSOP" in s:
        return "ttsop"
    if "\\SOP8\\" in s:
        return "sop8"
    return ""


def _grid_kind(ws: Any, ttsop_cell: str, sop8_cell: str, xlsx: Path | None) -> str:
    t_ok = _cell_ok(ws, ttsop_cell)
    s_ok = _cell_ok(ws, sop8_cell)
    t_n = _cell_num(ws, ttsop_cell)
    s_n = _cell_num(ws, sop8_cell)
    if t_n and not s_n:
        return "ttsop"
    if s_n and not t_n:
        return "sop8"
    if t_ok and not s_ok:
        return "ttsop"
    if s_ok and not t_ok:
        return "sop8"
    return _path_kind(xlsx)


def _xlsx_from_map(path: Path, data: dict[str, Any]) -> Path | None:
    rel = ""
    wb = data.get("workbook")
    if isinstance(wb, dict):
        rel = str(wb.get("path") or "")
    if rel:
        cand = (path.parent / rel).resolve()
        if cand.is_file():
            return cand
    folder = path.parent.parent / "workbook"
    if folder.is_dir():
        hits = sorted(
            x
            for x in folder.glob("*.xlsx")
            if not x.name.startswith("~$") and "AutoRecovered" not in x.name
        )
        if hits:
            return hits[0]
    return None


def outline_test(
    *,
    folder: str,
    excel_sheet: str,
    fixture_mode: str,
    sample: int,
    automated: bool = True,
) -> dict[str, Any]:
    return {
        "folder": folder,
        "excel_sheet": excel_sheet,
        "fixture_mode": str(fixture_mode or "LOGIC"),
        "automated": bool(automated),
        "dut_iterations": max(1, int(sample or 4)),
    }


def campaign_header(
    *,
    component: str,
    part: str,
    package: str,
    version: str,
    sample_size: int,
    workbook_name: str,
    operator: str = "",
) -> dict[str, Any]:
    out: dict[str, Any] = {
        "component": component,
        "part": part,
        "package": package,
        "version": version or "Version_1",
        "sample_size": max(1, int(sample_size or 4)),
        "workbook": {"path": f"../workbook/{workbook_name}"},
        "naming": dict(_NAMING),
        "tests": {},
    }
    if operator:
        out["operator"] = operator
    return out


def _clean_paste(entry: dict[str, Any]) -> None:
    paste = entry.get("paste")
    if paste is None:
        return
    if _is_fill(paste):
        entry.pop("paste", None)
        return
    if not isinstance(paste, dict):
        return
    cleaned = {k: v for k, v in paste.items() if not _is_fill(v)}
    if cleaned:
        entry["paste"] = cleaned
    else:
        entry.pop("paste", None)


def _merge_values(entry: dict[str, Any], values: dict[str, Any], excel_sheet: str | None = None) -> None:
    if excel_sheet:
        entry["excel_sheet"] = excel_sheet
    paste = entry.setdefault("paste", {})
    if not isinstance(paste, dict):
        paste = {}
        entry["paste"] = paste
    cur = paste.get("values")
    if not isinstance(cur, dict):
        cur = {}
        paste["values"] = cur
    for mid, cells in values.items():
        if mid not in cur:
            cur[mid] = cells


def _put_values(entry: dict[str, Any], values: dict[str, Any], excel_sheet: str | None = None) -> None:
    if excel_sheet:
        entry["excel_sheet"] = excel_sheet
    paste = entry.setdefault("paste", {})
    if not isinstance(paste, dict):
        paste = {}
        entry["paste"] = paste
    cur = paste.get("values")
    if not isinstance(cur, dict):
        cur = {}
        paste["values"] = cur
    for mid, cells in values.items():
        cur[mid] = cells


def _drop_ids(entry: dict[str, Any], ids: tuple[str, ...]) -> None:
    paste = entry.get("paste")
    if not isinstance(paste, dict):
        return
    cur = paste.get("values")
    if not isinstance(cur, dict):
        return
    for mid in ids:
        cur.pop(mid, None)
    if not cur:
        paste.pop("values", None)
    if not paste:
        entry.pop("paste", None)


def _pick_entry(tests: dict[str, Any], *names: str) -> dict[str, Any] | None:
    want = set()
    for n in names:
        want |= _fold_name(n)
    for key, entry in tests.items():
        if not isinstance(entry, dict):
            continue
        have = _fold_name(key) | _fold_name(str(entry.get("folder") or "")) | _fold_name(
            str(entry.get("excel_sheet") or "")
        )
        if want & have:
            return entry
    return None


def _fold_tok(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(name or "").lower())


def _sheet_by_fold(wb: Any, *names: str) -> str:
    folded = {_fold_tok(n): n for n in getattr(wb, "sheetnames", [])}
    for name in names:
        hit = folded.get(_fold_tok(name))
        if hit:
            return hit
    return ""


def _is_named_stub(ws: Any) -> bool:
    a1 = str(ws["A1"].value or "").strip().lower()
    return "(stub)" in a1 and not _probe_dut_value_cells(ws)


def attach_known_values(tests: dict[str, Any], sheets: list[str], xlsx: Path | None) -> None:
    sheet_set = {str(s) for s in sheets}
    if "Iplus" in sheet_set and (xlsx is None or not xlsx.is_file()):
        ip = _pick_entry(tests, "iplus")
        if ip is not None:
            _merge_values(ip, _IPLUS_STUB, "Iplus")
    if xlsx is None or not xlsx.is_file():
        return
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, data_only=False)
    dirty = False
    try:
        voh = _pick_entry(tests, "voh", "voh_load")
        vol = _pick_entry(tests, "vol", "vol_load")
        if "VOX" in sheet_set and "VOX" in wb.sheetnames:
            probed = _probe_vox_corners(wb["VOX"])
            voh_ids = {k: v for k, v in probed.items() if k.startswith("VOH_")}
            vol_ids = {k: v for k, v in probed.items() if k.startswith("VOL_")}
            if voh is not None:
                stale = tuple(
                    k
                    for k in ((voh.get("paste") or {}).get("values") or {})
                    if str(k).startswith("VOH_") and k not in voh_ids
                )
                if stale:
                    _drop_ids(voh, stale)
                if voh_ids:
                    _put_values(voh, voh_ids, "VOX")
            if vol is not None:
                stale = tuple(
                    k
                    for k in ((vol.get("paste") or {}).get("values") or {})
                    if str(k).startswith("VOL_") and k not in vol_ids
                )
                if stale:
                    _drop_ids(vol, stale)
                if vol_ids:
                    _put_values(vol, vol_ids, "VOX")
        icc_sheet = _sheet_by_fold(wb, "ICC", "IDD")
        icc = _pick_entry(tests, "supply_current", "supply_current_sweep", "icc", "idd")
        # RS0204 Icc is a condition grid -- do not stamp D10 onto it.
        if icc is not None and str(icc.get("excel_sheet") or "") != "Icc":
            probed_icc = _probe_max_ua(wb[icc_sheet], "ICC_uA") if icc_sheet else {}
            if probed_icc:
                _put_values(icc, probed_icc, icc_sheet)
            else:
                _drop_ids(icc, ("ICC_uA",))
        ioff_sheet = _sheet_by_fold(wb, "IOFF")
        if ioff_sheet:
            ioff = _pick_entry(tests, "ioff_leakage", "ioff")
            probed_ioff = _probe_max_ua(wb[ioff_sheet], "IOFF_uA")
            if ioff is not None:
                if probed_ioff:
                    _put_values(ioff, probed_ioff, ioff_sheet)
                else:
                    _drop_ids(ioff, ("IOFF_uA",))
        cin_sheet = _sheet_by_fold(wb, "CIN")
        if cin_sheet:
            cin = _pick_entry(tests, "cin")
            cin_cell = _probe_col_first_value(wb[cin_sheet], "average", "cin") or _probe_col_first_value(
                wb[cin_sheet], "cin", "pf"
            )
            if cin is not None and cin_cell:
                _put_values(cin, {"CIN_pF": cin_cell}, cin_sheet)
        cpd_sheet = _sheet_by_fold(wb, "CPD")
        if cpd_sheet:
            cpd = _pick_entry(tests, "cpd")
            cpd_cell = _probe_col_first_value(
                wb[cpd_sheet], "average", "cpd"
            ) or _probe_col_first_value(wb[cpd_sheet], "cpd", "pf")
            if cpd is not None and cpd_cell:
                _put_values(cpd, {"CPD_pF": cpd_cell}, cpd_sheet)
        for sheet, mid in _A1_IDS.items():
            if sheet not in sheet_set or sheet not in wb.sheetnames:
                continue
            entry = _pick_entry(tests, sheet, mid)
            cell = _probe_a1_cell(wb[sheet])
            if entry is not None and cell:
                _put_values(entry, {mid: cell}, sheet)
        if "GBW" in sheet_set and "GBW" in wb.sheetnames:
            gbw = _pick_entry(tests, "gbw")
            kind = _grid_kind(wb["GBW"], "R20", "C21", xlsx)
            if kind == "ttsop" and gbw is not None:
                _merge_values(gbw, _GBW_TTSOP)
            elif kind == "sop8" and gbw is not None:
                _merge_values(gbw, _GBW_SOP8)
        if "VOS" in sheet_set and "VOS" in wb.sheetnames:
            vos = _pick_entry(tests, "vos", "vos_sweep")
            kind = _grid_kind(wb["VOS"], "R16", "B16", xlsx)
            if kind == "ttsop" and vos is not None:
                _merge_values(vos, _VOS_TTSOP)
            elif kind == "sop8" and vos is not None:
                _merge_values(vos, _VOS_SOP8)
        ip = _pick_entry(tests, "iplus")
        ip_sheet = _sheet_by_fold(wb, "Iplus")
        if ip is not None and ip_sheet:
            ws_ip = wb[ip_sheet]
            dut_cells = _probe_dut_value_cells(ws_ip)
            if dut_cells:
                _put_values(ip, {"IPLUS_uA": dut_cells}, ip_sheet)
            else:
                _merge_values(ip, _IPLUS_STUB, ip_sheet)
                if _label_stub_mid(ws_ip, "IPLUS_uA"):
                    dirty = True
        for key, entry in list(tests.items()):
            if not isinstance(entry, dict):
                continue
            sheet = str(entry.get("excel_sheet") or entry.get("folder") or key)
            real = _sheet_by_fold(wb, sheet)
            if not real:
                continue
            ws = wb[real]
            if _is_named_stub(ws):
                mids = _mids_for_entry(str(key), entry)
                if mids:
                    _put_values(entry, {mids[0]: "B2"}, real)
                    if _label_stub_mid(ws, mids[0]):
                        dirty = True
                continue
            named = _probe_named_dut_mids(ws)
            mids = _mids_for_entry(str(key), entry)
            if named:
                keep = {k: v for k, v in named.items() if not mids or k in mids}
                _put_values(entry, keep or named, real)
                continue
            base = _probe_dut_value_cells(ws)
            if not mids or not base:
                continue
            payload = {mid: _shift_row(base, i) for i, mid in enumerate(mids)}
            _put_values(entry, payload, real)
        attach_known_photos(tests, sheets, wb=wb)
        attach_known_narrative(tests, sheets, wb=wb)
        if dirty:
            wb.save(xlsx)
    finally:
        wb.close()


def tests_from_sheets(sheets: list[str], *, family: str, sample: int) -> dict[str, Any]:
    tests: dict[str, Any] = {}
    used: set[str] = set()
    have = {str(s) for s in sheets}
    for name in sheets:
        if name in _SKIP_SHEETS:
            continue
        if name in ("VIX", "VIHL"):
            key = "VIH_VIL"
            if key not in used:
                used.add(key)
                tests[key] = outline_test(
                    folder=key,
                    excel_sheet=name,
                    fixture_mode=_infer_mode(key, family),
                    sample=sample,
                )
            continue
        if name == "VOX":
            for key in ("VOH", "VOL"):
                if key in used:
                    continue
                used.add(key)
                tests[key] = outline_test(
                    folder=key,
                    excel_sheet="VOX",
                    fixture_mode=_infer_mode(key, family),
                    sample=sample,
                )
            continue
        if name in ("VOH", "VOL") and "VOX" in have:
            continue
        if name in ("ICC", "IDD"):
            key = "Supply_Current"
            if key not in used:
                used.add(key)
                tests[key] = outline_test(
                    folder=key,
                    excel_sheet=name,
                    fixture_mode=_infer_mode(key, family),
                    sample=sample,
                )
            continue
        key = _folder_key(name)
        base = key
        n = 2
        while key in used:
            key = f"{base}{n}"
            n += 1
        used.add(key)
        tests[key] = outline_test(
            folder=key,
            excel_sheet=name,
            fixture_mode=_infer_mode(key, family),
            sample=sample,
            automated=True,
        )
    return tests


def outline_from_workbook(
    *,
    component: str,
    part: str,
    package: str,
    version: str,
    sample_size: int,
    workbook_name: str,
    sheets: list[str],
    family: str = "",
    operator: str = "",
    xlsx: Path | None = None,
) -> dict[str, Any]:
    fam = family or _family_from_component(component)
    data = campaign_header(
        component=component,
        part=part,
        package=package,
        version=version,
        sample_size=sample_size,
        workbook_name=workbook_name,
        operator=operator,
    )
    data["tests"] = tests_from_sheets(sheets, family=fam, sample=int(sample_size or 4))
    attach_known_values(data["tests"], sheets, xlsx)
    return data


def upgrade_sheet_map(
    data: dict[str, Any],
    *,
    family: str = "",
    sample: int = 4,
    sheets: list[str] | None = None,
    xlsx: Path | None = None,
) -> dict[str, Any]:
    out = deepcopy(data)
    fam = family or _family_from_component(str(out.get("component") or ""))
    out.setdefault("sample_size", max(1, int(sample or 4)))
    out.setdefault("naming", dict(_NAMING))
    if not isinstance(out.get("workbook"), dict) or not (out.get("workbook") or {}).get("path"):
        name = xlsx.name if xlsx is not None else f"{out.get('part') or 'Lab'}_Lab_Report.xlsx"
        out["workbook"] = {"path": f"../workbook/{name}"}
    if isinstance(out.get("naming"), dict):
        nam = dict(out["naming"])
        nam.setdefault("screenshot", _NAMING["screenshot"])
        nam.setdefault("graph", _NAMING["graph"])
        out["naming"] = nam
    tests = out.get("tests")
    if not isinstance(tests, dict) or not tests:
        if sheets:
            out["tests"] = tests_from_sheets(sheets, family=fam, sample=int(out["sample_size"]))
            tests = out["tests"]
        else:
            out["tests"] = {}
            tests = out["tests"]
    sample_n = int(out.get("sample_size") or sample or 4)
    for key, entry in list(tests.items()):
        if not isinstance(entry, dict):
            tests.pop(key, None)
            continue
        sheet = str(entry.get("excel_sheet") or "")
        folder = str(entry.get("folder") or key)
        if sheet in _SKIP_SHEETS or folder in _SKIP_SHEETS or key in _SKIP_SHEETS:
            tests.pop(key, None)
            continue
        entry.setdefault("folder", key)
        entry.setdefault("excel_sheet", str(entry.get("folder") or key))
        mode = str(entry.get("fixture_mode") or "")
        was_fill = (not mode) or mode.upper() in _FILL
        if was_fill:
            entry["fixture_mode"] = _infer_mode(str(entry.get("folder") or key), fam)
        if "automated" not in entry or was_fill:
            entry["automated"] = True
        entry.setdefault("dut_iterations", sample_n)
        _clean_paste(entry)
    attach_known_values(tests, list(sheets or []), xlsx)
    have = {str(s) for s in (sheets or [])}
    if have:
        for key, entry in list(tests.items()):
            if not isinstance(entry, dict):
                continue
            sheet = str(entry.get("excel_sheet") or "")
            if sheet and sheet not in have:
                tests.pop(key, None)
    return out


def iter_campaign_maps(base: Path | None = None) -> list[Path]:
    root = Path(base) if base else TEST_DB_ROOT
    if not root.is_dir():
        return []
    out: list[Path] = []
    for p in root.rglob("sheet_map.yaml"):
        if p.parent.name != "_manifest":
            continue
        if any(part.startswith("_") and part not in {"_manifest"} for part in p.parts):
            continue
        out.append(p)
    return sorted(out)


def outline_errors(data: dict[str, Any]) -> list[str]:
    errors: list[str] = []
    for key in ("component", "part", "package", "version", "sample_size", "tests"):
        if key not in data or data.get(key) in (None, ""):
            errors.append(f"missing {key}")
    nam = data.get("naming")
    if not isinstance(nam, dict) or not nam.get("screenshot") or not nam.get("graph"):
        errors.append("naming.screenshot/graph missing")
    wb = data.get("workbook")
    if not isinstance(wb, dict) or not wb.get("path"):
        errors.append("workbook.path missing")
    tests = data.get("tests")
    if not isinstance(tests, dict):
        errors.append("tests must be a map")
        return errors
    for key, entry in tests.items():
        if not isinstance(entry, dict):
            errors.append(f"{key}: not a map")
            continue
        for field in TEST_OUTLINE:
            if field not in entry or entry.get(field) in (None, ""):
                errors.append(f"{key}: missing {field}")
        mode = str(entry.get("fixture_mode") or "")
        if mode.upper() in _FILL:
            errors.append(f"{key}: fixture_mode FILL_ME")
        paste = entry.get("paste")
        if isinstance(paste, dict):
            vals = paste.get("values")
            if _is_fill(vals) or (isinstance(vals, dict) and any(_is_fill(v) for v in vals.values())):
                errors.append(f"{key}: paste.values still FILL_ME")
            if _is_fill(paste.get("photos")):
                errors.append(f"{key}: paste.photos FILL_ME (omit until measured)")
            if _is_fill(paste.get("narrative")):
                errors.append(f"{key}: paste.narrative FILL_ME (omit until probed)")
        elif _is_fill(paste):
            errors.append(f"{key}: paste FILL_ME")
    return errors


def _merge_enabled_tests(work: dict[str, Any], *, family: str, sample: int) -> None:
    """Add missing enabled TestSpec lab_sheets. Does not invent paste cells."""
    from ate.core.new_product import part_key_for, suite_for_part, _test_track_rows
    from ate.core.registry import get as reg_get

    part = str(work.get("part") or "")
    package = str(work.get("package") or "")
    pk = part_key_for(part)
    if not pk:
        return
    fam = family or suite_for_part(
        part, component=str(work.get("component") or ""), package=package
    )
    tests = work.get("tests")
    if not isinstance(tests, dict):
        tests = {}
        work["tests"] = tests
    for tid, folder in _test_track_rows(pk, fam):
        if _pick_entry(tests, tid, folder) is not None:
            continue
        spec = reg_get(tid)
        mode = str(getattr(spec, "fixture_mode", "") or "") or _infer_mode(tid, fam)
        tests[tid] = outline_test(
            folder=folder or tid,
            excel_sheet=folder or tid,
            fixture_mode=mode,
            sample=int(sample or 4),
        )


def apply_campaign_map(path: Path, *, xlsx: Path | None = None) -> dict[str, Any]:
    loaded = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    if not isinstance(loaded, dict):
        loaded = {}
    work = deepcopy(loaded)
    try:
        work.setdefault("component", path.parents[5].name)
        work.setdefault("part", path.parents[4].name)
        work.setdefault("package", path.parents[3].name)
        work.setdefault("version", path.parents[1].name)
    except IndexError:
        pass
    xlsx = Path(xlsx) if xlsx else _xlsx_from_map(path, work)
    sheets: list[str] = []
    if xlsx is not None:
        from openpyxl import load_workbook

        try:
            wb = load_workbook(xlsx, read_only=True)
            try:
                sheets = list(wb.sheetnames)
            finally:
                wb.close()
        except Exception:
            sheets = []
            xlsx = None
    if xlsx is not None:
        want = f"../workbook/{xlsx.name}"
        wb_block = work.get("workbook")
        if not isinstance(wb_block, dict):
            work["workbook"] = {"path": want}
        elif wb_block.get("path") != want:
            wb_block["path"] = want
    try:
        fam = _family_from_component(str(work.get("component") or path.parents[5].name))
    except IndexError:
        fam = _family_from_component(str(work.get("component") or ""))
    sample = int(work.get("sample_size") or 4)
    _merge_enabled_tests(work, family=fam, sample=sample)
    upgraded = upgrade_sheet_map(work, family=fam, sample=sample, sheets=sheets, xlsx=xlsx)
    if upgraded == loaded:
        return {"path": str(path), "wrote": False}
    path.write_text(yaml.safe_dump(upgraded, sort_keys=False, allow_unicode=True), encoding="utf-8")
    return {"path": str(path), "wrote": True}


def apply_all(base: Path | None = None) -> dict[str, Any]:
    wrote = 0
    kept = 0
    for p in iter_campaign_maps(base):
        row = apply_campaign_map(p)
        if row.get("wrote"):
            wrote += 1
        else:
            kept += 1
    return {"ok": True, "wrote": wrote, "kept": kept, "count": wrote + kept}


def main(argv: list[str] | None = None) -> int:
    import sys

    args = list(sys.argv[1:] if argv is None else argv)
    if "--apply" in args:
        res = apply_all()
        print(f"OK campaign_outline apply wrote={res['wrote']} kept={res['kept']} count={res['count']}")
        return 0
    from ate.core.check_campaign_outline import main as check_main

    return check_main()


if __name__ == "__main__":
    raise SystemExit(main())
