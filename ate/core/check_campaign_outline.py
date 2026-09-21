"""Fail-closed RS622 campaign outline. FILL_ME and missing fixture_mode must fail.

Run: python -m ate.core.check_campaign_outline
"""
from __future__ import annotations

import tempfile
from pathlib import Path

import yaml
from openpyxl import Workbook

from ate.core.campaign_outline import (
    attach_known_values,
    iter_campaign_maps,
    outline_errors,
    outline_from_workbook,
    outline_test,
    tests_from_sheets,
    upgrade_sheet_map,
    _xlsx_from_map,
    _SKIP_SHEETS,
)
from ate.core.paths import PART_DB_ROOT, TEST_DB_ROOT


def _header(**extra: object) -> dict:
    data = {
        "component": "Logic",
        "part": "X",
        "package": "Y",
        "version": "Version_1",
        "sample_size": 4,
        "workbook": {"path": "../workbook/x.xlsx"},
        "naming": {
            "screenshot": "{TEST}_{DUT}_{VARIANT}_{TIMESTAMP}.jpg",
            "graph": "{TEST}_{DUT}_{VARIANT}_{TIMESTAMP}.png",
        },
        "tests": {},
    }
    data.update(extra)
    return data


def _has_fill_token(obj: object) -> bool:
    """Exact FILL_ME values only. Comments like 'Do not write FILL_ME.' must pass."""
    if obj == "FILL_ME":
        return True
    if isinstance(obj, dict):
        return any(_has_fill_token(k) or _has_fill_token(v) for k, v in obj.items())
    if isinstance(obj, list):
        return any(_has_fill_token(v) for v in obj)
    return False


def _memory_errors() -> list[str]:
    errors: list[str] = []
    missing = _header(
        tests={"VOH": {"folder": "VOH", "excel_sheet": "VOH"}},
    )
    miss = outline_errors(missing)
    if not any("fixture_mode" in e for e in miss):
        errors.append("outline_errors must fail when fixture_mode is missing")

    fill = _header(
        tests={
            "GBW": {
                "folder": "GBW",
                "excel_sheet": "GBW",
                "fixture_mode": "FILL_ME",
                "automated": False,
                "dut_iterations": 4,
                "paste": {"values": "FILL_ME", "photos": "FILL_ME"},
            },
            "Summary": {
                "folder": "Summary",
                "excel_sheet": "Summary",
                "fixture_mode": "FILL_ME",
                "automated": False,
                "dut_iterations": 4,
            },
        }
    )
    if not outline_errors(fill):
        errors.append("FILL_ME fixture_mode/paste must fail outline_errors")
    upgraded = upgrade_sheet_map(fill, family="opamp", sheets=["GBW", "Summary"])
    dumped = yaml.safe_dump(upgraded)
    if "FILL_ME" in dumped:
        errors.append("upgrade_sheet_map must strip FILL_ME")
    if "Summary" in (upgraded.get("tests") or {}):
        errors.append("upgrade must drop Summary stub")
    gbw = (upgraded.get("tests") or {}).get("GBW") or {}
    if gbw.get("fixture_mode") != "G11":
        errors.append(f"GBW FILL_ME must infer G11, got {gbw.get('fixture_mode')!r}")
    if gbw.get("automated") is not True:
        errors.append("FILL_ME stub automated false must become true")
    if not _has_fill_token(fill):
        errors.append("_has_fill_token must catch exact FILL_ME values")
    comment_doc = yaml.safe_load("# Do not write FILL_ME.\npart: RS1G07\n")
    if _has_fill_token(comment_doc):
        errors.append("FILL_ME in a YAML comment must not count as a value")

    from_sheets = outline_from_workbook(
        component="Logic",
        part="RS1G08",
        package="SOT23",
        version="Version_1",
        sample_size=4,
        workbook_name="RS1G08_Lab_Report.xlsx",
        sheets=["Summary", "VOX", "ICC", "VOH"],
        family="logic",
    )
    tests = from_sheets.get("tests") or {}
    if "Summary" in tests:
        errors.append("tests_from_sheets must skip Summary")
    voh = tests.get("VOH") or {}
    if voh.get("excel_sheet") != "VOX":
        errors.append(f"VOX workbook must map VOH excel_sheet VOX, got {voh}")
    if (tests.get("Supply_Current") or {}).get("excel_sheet") != "ICC":
        errors.append("ICC sheet must map Supply_Current excel_sheet ICC")
    idd_tests = tests_from_sheets(["IDD"], family="logic", sample=4)
    if (idd_tests.get("Supply_Current") or {}).get("excel_sheet") != "IDD":
        errors.append("IDD sheet must map Supply_Current excel_sheet IDD")
    if outline_errors(from_sheets):
        errors.append("outline_from_workbook result must pass outline_errors: " + str(outline_errors(from_sheets)))

    tmp = Path(tempfile.mkdtemp(prefix="ate_outline_"))
    xlsx = tmp / "lab.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "VOX"
    wb.create_sheet("ICC")
    wb.save(xlsx)
    wb.close()
    hooked = {
        "voh_load": outline_test(
            folder="VOH",
            excel_sheet="VOH",
            fixture_mode="LOGIC",
            sample=4,
        ),
        "vol_load": outline_test(
            folder="VOL",
            excel_sheet="VOL",
            fixture_mode="LOGIC",
            sample=4,
        ),
        "supply_current_sweep": outline_test(
            folder="Supply_Current",
            excel_sheet="Supply_Current",
            fixture_mode="LOGIC",
            sample=4,
        ),
    }
    attach_known_values(hooked, ["VOX", "ICC"], xlsx)
    if hooked["voh_load"].get("paste", {}).get("values", {}).get("VOH_4p5V"):
        errors.append("empty VOX must not invent VOH_4p5V G16")
    if hooked["supply_current_sweep"].get("paste", {}).get("values", {}).get("ICC_uA"):
        errors.append("empty ICC must not invent ICC_uA D10")

    wb = Workbook()
    ws = wb.active
    ws.title = "VOX"
    ws["A10"] = "VOH"
    ws["A11"] = "Vcc (V)"
    ws["A12"] = 1.65
    ws["A16"] = 4.5
    ws["A19"] = "VOL"
    ws["A20"] = "Vcc (V)"
    ws["A21"] = 1.65
    ws["A25"] = 4.5
    icc = wb.create_sheet("ICC")
    icc["B6"] = "When input A is 0V, maximum ICC is 0.004uA, minimum ICC is 0.001uA."
    icc["C10"] = "maximum (uA)"
    icc["D10"] = 1
    wb.save(xlsx)
    wb.close()
    attach_known_values(hooked, ["VOX", "ICC"], xlsx)
    if hooked["voh_load"].get("excel_sheet") != "VOX":
        errors.append("attach_known_values must remap voh_load to VOX")
    if hooked["voh_load"].get("paste", {}).get("values", {}).get("VOH_4p5V") != ["G16", "H16", "I16"]:
        errors.append("RS1G08-shaped VOX must set VOH_4p5V G16:I16")
    if hooked["voh_load"].get("paste", {}).get("values", {}).get("VOH_1p65V") != ["G12", "H12", "I12"]:
        errors.append("RS1G08-shaped VOX must set VOH_1p65V G12:I12")
    if hooked["vol_load"].get("paste", {}).get("values", {}).get("VOL_4p5V") != ["G25", "H25", "I25"]:
        errors.append("RS1G08-shaped VOX must set VOL_4p5V G25:I25")
    if hooked["vol_load"].get("paste", {}).get("values", {}).get("VOL_1p65V") != ["G21", "H21", "I21"]:
        errors.append("RS1G08-shaped VOX must set VOL_1p65V G21:I21")
    if hooked["supply_current_sweep"].get("paste", {}).get("values", {}).get("ICC_uA") != "D10":
        errors.append("ICC maximum banner must set ICC_uA D10")

    wb = Workbook()
    ws = wb.active
    ws.title = "VOX"
    ws["A10"] = "GBW (Average)"
    ws["A15"] = "VOH"
    ws["A16"] = "Vcc (V)"
    ws["A17"] = 1.65
    ws["A21"] = 4.5
    ws["A24"] = "VOL"
    ws["A25"] = "Vcc (V)"
    ws["A26"] = 1.65
    ws["A30"] = 4.5
    wb.create_sheet("ICC")
    wb.save(xlsx)
    wb.close()
    attach_known_values(hooked, ["VOX", "ICC"], xlsx)
    if hooked["voh_load"].get("paste", {}).get("values", {}).get("VOH_4p5V") != ["G21", "H21", "I21"]:
        errors.append("shifted VOX must remap VOH_4p5V to G21, not keep G16")
    if hooked["voh_load"].get("paste", {}).get("values", {}).get("VOH_1p65V") != ["G17", "H17", "I17"]:
        errors.append("shifted VOX must set VOH_1p65V G17")
    if hooked["vol_load"].get("paste", {}).get("values", {}).get("VOL_4p5V") != ["G30", "H30", "I30"]:
        errors.append("shifted VOX must remap VOL_4p5V to G30")
    if hooked["vol_load"].get("paste", {}).get("values", {}).get("VOL_1p65V") != ["G26", "H26", "I26"]:
        errors.append("shifted VOX must set VOL_1p65V G26")
    if hooked["supply_current_sweep"].get("paste", {}).get("values", {}).get("ICC_uA"):
        errors.append("ICC without maximum banner must drop ICC_uA")

    from ate.core.campaign_outline import attach_known_photos

    photo_wb = Workbook()
    pws = photo_wb.active
    pws.title = "GBW"
    pws["A5"] = "Test Conclusion"
    pws["A8"] = "GBW #1"
    pws.merge_cells("A10:D19")
    pws.merge_cells("E10:H19")
    tests_photos = {
        "GBW": {
            "folder": "GBW",
            "excel_sheet": "GBW",
            "fixture_mode": "G11",
            "automated": True,
            "dut_iterations": 2,
            "paste": {
                "photos": {
                    "u1_chA": "A70",
                    "u1_chB": "E70",
                    "u2_chA": "I70",
                    "u2_chB": "M70",
                    "u3_chA": "Q70",
                    "u3_chB": "U70",
                    "u4_chA": "Y70",
                    "u4_chB": "AC70",
                }
            },
        }
    }
    attach_known_photos(tests_photos, ["GBW"], wb=photo_wb)
    got = (tests_photos["GBW"].get("paste") or {}).get("photos") or {}
    if got.get("u1_chA") != "A10" or got.get("u1_chB") != "E10":
        errors.append(f"attach_known_photos must discover merged boxes, got {got}")
    if "u4_chA" in got:
        errors.append("must not keep invented DUT4 photo cells")

    from ate.core.campaign_outline import attach_known_narrative

    nws = photo_wb.create_sheet("Slew Rate")
    nws["A3"] = "Test Conditions"
    nws["B3"] = "#VALUE!"
    nws["A8"] = "Datasheet"
    nws["B8"] = None
    nws["A15"] = "Test Conclusion"
    nws["B15"] = None
    tests_nar = {
        "SlewRate": {
            "folder": "Slew Rate",
            "excel_sheet": "Slew Rate",
            "fixture_mode": "BUFFER",
            "automated": True,
            "dut_iterations": 4,
        }
    }
    attach_known_narrative(tests_nar, ["Slew Rate"], wb=photo_wb)
    nar = (tests_nar["SlewRate"].get("paste") or {}).get("narrative") or {}
    if nar.get("conditions") != "B3" or nar.get("datasheet") != "B8" or nar.get("conclusion") != "B15":
        errors.append(f"attach_known_narrative must probe intro B cells, got {nar}")
    photo_wb.close()
    return errors


def _switch_dut_grid_errors() -> list[str]:
    """Parameter DUT_1 headers must not keep Iplus B2 (that cell is Unit)."""
    errors: list[str] = []
    tmp = Path(tempfile.mkdtemp(prefix="ate_sw_")) / "sw.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Iplus"
    ws["A1"] = "Parameter"
    ws["B1"] = "Unit"
    ws["F1"] = "DUT_1"
    ws["G1"] = "DUT_2"
    ws["H1"] = "DUT_3"
    ws["I1"] = "DUT_4"
    ws["A2"] = "(empty -- run START to fill)"
    off = wb.create_sheet("LeakageOff")
    off["A1"] = "Parameter"
    off["F1"] = "DUT_1"
    off["G1"] = "DUT_2"
    wb.save(tmp)
    tests = {
        "iplus": outline_test(
            folder="Iplus", excel_sheet="Iplus", fixture_mode="LIM_RS2323", sample=4
        ),
        "leakage_off": outline_test(
            folder="LeakageOff", excel_sheet="LeakageOff", fixture_mode="LIM_RS2323", sample=4
        ),
    }
    attach_known_values(tests, ["Iplus", "LeakageOff"], tmp)
    ip_cells = ((tests["iplus"].get("paste") or {}).get("values") or {}).get("IPLUS_uA")
    if ip_cells != ["F2", "G2", "H2", "I2"]:
        errors.append(f"Parameter Iplus must probe DUT F2:I2, got {ip_cells}")
    ioz = ((tests["leakage_off"].get("paste") or {}).get("values") or {}).get("IOZ_uA")
    if ioz != ["F2", "G2"]:
        errors.append(f"Parameter LeakageOff must probe DUT F2:G2, got {ioz}")

    stub_path = Path(tempfile.mkdtemp(prefix="ate_ip_")) / "stub.xlsx"
    wb2 = Workbook()
    ip = wb2.active
    ip.title = "Iplus"
    ip["A1"] = "Iplus (stub)"
    wb2.save(stub_path)
    stub_tests = {
        "iplus": outline_test(
            folder="Iplus", excel_sheet="Iplus", fixture_mode="LIM_RS2323", sample=4
        )
    }
    attach_known_values(stub_tests, ["Iplus"], stub_path)
    stub_cell = ((stub_tests["iplus"].get("paste") or {}).get("values") or {}).get("IPLUS_uA")
    if stub_cell != "B2":
        errors.append(f"SeeLim-style Iplus stub must stay B2, got {stub_cell}")

    ton_path = Path(tempfile.mkdtemp(prefix="ate_ton_")) / "ton.xlsx"
    wb3 = Workbook()
    ton = wb3.active
    ton.title = "TonToff"
    ton["A1"] = "Parameter"
    ton["F1"] = "DUT_1"
    ton["G1"] = "DUT_2"
    ton["A2"] = "USB_TON_ns"
    ton["A3"] = "USB_TOFF_ns"
    wb3.save(ton_path)
    ton_tests = {
        "usb_ton_toff": outline_test(
            folder="TonToff", excel_sheet="TonToff", fixture_mode="LIM_RS2323", sample=2
        )
    }
    attach_known_values(ton_tests, ["TonToff"], ton_path)
    ton_vals = ((ton_tests["usb_ton_toff"].get("paste") or {}).get("values") or {})
    if ton_vals.get("USB_TON_ns") != ["F2", "G2"]:
        errors.append(f"USB_TON_ns must probe F2:G2, got {ton_vals.get('USB_TON_ns')}")
    if ton_vals.get("USB_TOFF_ns") != ["F3", "G3"]:
        errors.append(f"USB_TOFF_ns must probe F3:G3, got {ton_vals.get('USB_TOFF_ns')}")

    cin_path = Path(tempfile.mkdtemp(prefix="ate_cin_")) / "cin.xlsx"
    wb4 = Workbook()
    cin = wb4.active
    cin.title = "CIN"
    cin["A1"] = "Parameter"
    cin["F1"] = "DUT_1"
    cin["G1"] = "DUT_2"
    cin["H1"] = "DUT_3"
    cin["I1"] = "DUT_4"
    cin["A2"] = "CIN_pF"
    tp = wb4.create_sheet("TP")
    tp["A1"] = "Parameter"
    tp["F1"] = "DUT_1"
    tp["G1"] = "DUT_2"
    tp["A2"] = "CLKQ_PHL_ns"
    tp["A3"] = "CLKQ_PLH_ns"
    wb4.save(cin_path)
    cin_tests = {
        "cin": outline_test(folder="CIN", excel_sheet="CIN", fixture_mode="LOGIC", sample=4),
        "clk_q": outline_test(folder="TP", excel_sheet="TP", fixture_mode="LOGIC", sample=2),
    }
    attach_known_values(cin_tests, ["CIN", "TP"], cin_path)
    cin_cells = ((cin_tests["cin"].get("paste") or {}).get("values") or {}).get("CIN_pF")
    if cin_cells != ["F2", "G2", "H2", "I2"]:
        errors.append(f"Parameter CIN must probe DUT F2:I2, got {cin_cells}")
    clk_vals = ((cin_tests["clk_q"].get("paste") or {}).get("values") or {})
    if clk_vals.get("CLKQ_PHL_ns") != ["F2", "G2"]:
        errors.append(f"CLKQ_PHL_ns must probe F2:G2, got {clk_vals.get('CLKQ_PHL_ns')}")
    if clk_vals.get("CLKQ_PLH_ns") != ["F3", "G3"]:
        errors.append(f"CLKQ_PLH_ns must probe F3:G3, got {clk_vals.get('CLKQ_PLH_ns')}")

    ariff_cin = Path(tempfile.mkdtemp(prefix="ate_acn_")) / "cin.xlsx"
    wb5 = Workbook()
    acn = wb5.active
    acn.title = "Cin"
    acn["A20"] = "Input A (V)"
    acn["H20"] = "Average Cin(pF)"
    acn["H21"] = 5.2
    cpd = wb5.create_sheet("Cpd")
    cpd["D8"] = "Average Cpd"
    cpd["E8"] = "=AVERAGE(D12,D17)"
    wb5.save(ariff_cin)
    acn_tests = {
        "cin": outline_test(folder="Cin", excel_sheet="Cin", fixture_mode="LOGIC", sample=3),
        "cpd": outline_test(folder="Cpd", excel_sheet="Cpd", fixture_mode="LOGIC", sample=3),
    }
    attach_known_values(acn_tests, ["Cin", "Cpd"], ariff_cin)
    if ((acn_tests["cin"].get("paste") or {}).get("values") or {}).get("CIN_pF") != "H21":
        errors.append(
            "Ariff Cin must probe Average Cin H21, got "
            + str(((acn_tests["cin"].get("paste") or {}).get("values") or {}).get("CIN_pF"))
        )
    if ((acn_tests["cpd"].get("paste") or {}).get("values") or {}).get("CPD_pF") != "E8":
        errors.append(
            "Ariff Cpd must probe Average Cpd E8, got "
            + str(((acn_tests["cpd"].get("paste") or {}).get("values") or {}).get("CPD_pF"))
        )

    leak_path = Path(tempfile.mkdtemp(prefix="ate_lk_")) / "lk.xlsx"
    wb6 = Workbook()
    off = wb6.active
    off.title = "LeakageOff"
    off["A1"] = "LeakageOff (stub)"
    wb6.save(leak_path)
    leak_tests = {
        "leakage_off": outline_test(
            folder="LeakageOff", excel_sheet="LeakageOff", fixture_mode="LIM_RS2323", sample=4
        )
    }
    attach_known_values(leak_tests, ["LeakageOff"], leak_path)
    leak_cell = ((leak_tests["leakage_off"].get("paste") or {}).get("values") or {}).get("IOZ_uA")
    if leak_cell != "B2":
        errors.append(f"SeeLim LeakageOff stub must probe B2, got {leak_cell}")
    return errors


def _live_errors() -> list[str]:
    errors: list[str] = []
    gold = PART_DB_ROOT / "_manifest" / "sheet_map.yaml"
    if not gold.is_file():
        return [f"RS622 TTSOP8 sheet_map missing: {gold}"]
    gold_data = yaml.safe_load(gold.read_text(encoding="utf-8")) or {}
    gbw_vals = (
        ((gold_data.get("tests") or {}).get("GBW") or {}).get("paste") or {}
    ).get("values") or {}
    if "GBW_MHz" not in gbw_vals:
        errors.append("RS622 TTSOP8 tests.GBW.paste.values.GBW_MHz missing")
    gold_errs = outline_errors(gold_data)
    if gold_errs:
        errors.append("RS622 TTSOP8 outline: " + "; ".join(gold_errs[:8]))

    logic = (
        TEST_DB_ROOT
        / "Logic"
        / "RS1G08"
        / "SOT23"
        / "Ariff"
        / "Version_1"
        / "_manifest"
        / "sheet_map.yaml"
    )
    if logic.is_file():
        logic_data = yaml.safe_load(logic.read_text(encoding="utf-8")) or {}
        tests = logic_data.get("tests") or {}
        voh = tests.get("VOH") or tests.get("voh_load") or {}
        vals = ((voh.get("paste") or {}).get("values") or {})
        if "VOH_4p5V" not in vals:
            errors.append("RS1G08 Ariff SOT23 must keep VOH_4p5V on VOX")
        xlsx = _xlsx_from_map(logic, logic_data)
        if xlsx is not None:
            from openpyxl import load_workbook

            wb = load_workbook(xlsx, read_only=True)
            try:
                names = set(wb.sheetnames)
            finally:
                wb.close()
            if "VOX" in names and str(voh.get("excel_sheet") or "") != "VOX":
                errors.append("RS1G08 workbook has VOX but excel_sheet is not VOX")
            if vals.get("VOH_4p5V") != ["G16", "H16", "I16"]:
                errors.append(f"RS1G08 VOX 4.5V DUT must stay G16, got {vals.get('VOH_4p5V')}")
            if vals.get("VOH_1p65V") != ["G12", "H12", "I12"]:
                errors.append(f"RS1G08 VOX 1.65V DUT must be G12, got {vals.get('VOH_1p65V')}")
    else:
        errors.append(f"RS1G08 Ariff sheet_map missing: {logic}")

    rs32 = (
        TEST_DB_ROOT
        / "Logic"
        / "RS1G32"
        / "SOT23"
        / "Ariff"
        / "Version_1"
        / "_manifest"
        / "sheet_map.yaml"
    )
    if rs32.is_file():
        rs32_data = yaml.safe_load(rs32.read_text(encoding="utf-8")) or {}
        rs32_voh = (rs32_data.get("tests") or {}).get("VOH") or {}
        rs32_cells = ((rs32_voh.get("paste") or {}).get("values") or {}).get("VOH_4p5V")
        if rs32_cells != ["G21", "H21", "I21"]:
            errors.append(f"RS1G32 VOX 4.5V DUT must be G21 (table shifted), got {rs32_cells}")
        rs32_165 = ((rs32_voh.get("paste") or {}).get("values") or {}).get("VOH_1p65V")
        if rs32_165 != ["G17", "H17", "I17"]:
            errors.append(f"RS1G32 VOX 1.65V DUT must be G17, got {rs32_165}")
        rs32_vol = ((rs32_data.get("tests") or {}).get("VOL") or {}).get("paste") or {}
        vol_cells = (rs32_vol.get("values") or {}).get("VOL_4p5V")
        if vol_cells == ["G25", "H25", "I25"]:
            errors.append("RS1G32 VOL_4p5V must not stay on RS1G08 G25")

    seelim_msop = (
        TEST_DB_ROOT
        / "AnalogSwitch"
        / "RS2323"
        / "MSOP"
        / "SeeLim"
        / "Version_1"
        / "_manifest"
        / "sheet_map.yaml"
    )
    if seelim_msop.is_file():
        sm = yaml.safe_load(seelim_msop.read_text(encoding="utf-8")) or {}
        wb_rel = str(((sm.get("workbook") or {}) if isinstance(sm, dict) else {}).get("path") or "")
        if "RS2323_Lab_Report_MSOP.xlsx" not in wb_rel:
            errors.append(f"SeeLim RS2323 MSOP map must point at package-named xlsx, got {wb_rel}")
        tests = sm.get("tests") if isinstance(sm.get("tests"), dict) else {}
        ioz = ((tests.get("leakage_off") or {}).get("paste") or {}).get("values") or {}
        if ioz.get("IOZ_uA") != "B2":
            errors.append(f"SeeLim RS2323 MSOP LeakageOff must stay B2, got {ioz.get('IOZ_uA')}")
        ron = ((tests.get("ron") or {}).get("paste") or {}).get("values") or {}
        if ron.get("RON_ohm") != ["F2", "G2", "H2", "I2"]:
            errors.append(f"SeeLim RS2323 MSOP Ron DUT must be F2:I2, got {ron.get('RON_ohm')}")
        xlsx = _xlsx_from_map(seelim_msop, sm)
        if xlsx is None or not xlsx.is_file():
            errors.append("SeeLim RS2323 MSOP workbook missing")
        else:
            from openpyxl import load_workbook

            wb = load_workbook(xlsx, read_only=True)
            try:
                names = set(wb.sheetnames)
            finally:
                wb.close()
            for need in ("Ron", "Vth", "TON", "CinConCoff", "TBBM", "LeakageOff"):
                if need not in names:
                    errors.append(f"SeeLim RS2323 MSOP missing sheet {need}")
    rs2227_map = (
        TEST_DB_ROOT
        / "AnalogSwitch"
        / "RS2227"
        / "MSOP"
        / "SeeLim"
        / "Version_1"
        / "_manifest"
        / "sheet_map.yaml"
    )
    if rs2227_map.is_file():
        sm = yaml.safe_load(rs2227_map.read_text(encoding="utf-8")) or {}
        xlsx = _xlsx_from_map(rs2227_map, sm)
        if xlsx is None or not xlsx.is_file():
            errors.append("SeeLim RS2227 MSOP workbook missing")
        usb = ((sm.get("tests") or {}).get("usb_ron") or {}).get("paste") or {}
        if (usb.get("values") or {}).get("USB_RON_ohm") != ["F2", "G2", "H2", "I2"]:
            errors.append(
                "SeeLim RS2227 usb_ron DUT must be F2:I2, got "
                + str((usb.get("values") or {}).get("USB_RON_ohm"))
            )

    for path in iter_campaign_maps():
        raw = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
        if not isinstance(raw, dict):
            errors.append(f"{path}: not a map")
            continue
        rel = str(path)
        for err in outline_errors(raw):
            errors.append(f"{rel}: {err}")
        if _has_fill_token(raw):
            errors.append(f"{rel}: FILL_ME still present")
        tests = raw.get("tests") if isinstance(raw.get("tests"), dict) else {}
        xlsx = _xlsx_from_map(path, raw)
        if not tests and xlsx is not None and xlsx.is_file():
            from openpyxl import load_workbook

            wb = load_workbook(xlsx, read_only=True)
            try:
                names = {str(n) for n in wb.sheetnames}
            finally:
                wb.close()
            meta = _SKIP_SHEETS | {"Setup", "Sweep", "Sheet1"}
            if names - meta:
                errors.append(f"{rel}: empty tests with workbook present")
        if len(errors) > 40:
            errors.append("... truncated campaign outline errors")
            break
    return errors


def main() -> int:
    errors = _memory_errors() + _switch_dut_grid_errors() + _live_errors()
    if errors:
        print("FAIL campaign_outline:")
        for line in errors:
            print(f"  - {line}")
        return 1
    print("OK campaign_outline: RS622 keys, no FILL_ME, known VOX/ICC/GBW paste")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
