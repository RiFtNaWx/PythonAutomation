"""Fail-closed: per-SKU golden refs, not one Parameter table.

Run: python -m ate.core.check_golden_refs

Proves: Parameter stubs are never the golden; RS1G08 senior book keeps VIX/VOX;
JianHong stub is not selected when a senior book exists.
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path

from openpyxl import Workbook

from ate.core.golden_refs import (
    _best_by_key,
    inspect_workbook,
    is_stub_workbook,
    resolve_golden,
)
from ate.core.paths import TEST_DB_ROOT


def _write_stub(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Device"
    ws["B1"] = "RS1G08"
    ws["B8"] = "Clean golden template -- no measured values. Fill via START / Fill Excel."
    vix = wb.create_sheet("VIH_VIL")
    vix["A1"] = "Parameter"
    vix["B1"] = "Unit"
    vox = wb.create_sheet("VOH")
    vox["A1"] = "Parameter"
    wb.save(path)
    wb.close()


def _write_senior(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Version Summary"
    vix = wb.create_sheet("VIX")
    vix["A1"] = "Test Parameter"
    vox = wb.create_sheet("VOX")
    vox["A1"] = "Test Parameter"
    wb.save(path)
    wb.close()


def _write_named_stub(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Summary (stub)"
    ip = wb.create_sheet("Iplus")
    ip["A1"] = "Iplus (stub)"
    off = wb.create_sheet("LeakageOff")
    off["A1"] = "LeakageOff (stub)"
    wb.save(path)
    wb.close()


def main() -> int:
    tmp = Path(tempfile.mkdtemp(prefix="ate_golden_check_"))
    stub = tmp / "stub.xlsx"
    senior = tmp / "ariff.xlsx"
    _write_stub(stub)
    _write_senior(senior)
    if not is_stub_workbook(stub):
        raise AssertionError("Parameter + Clean golden template must be stub")
    named = tmp / "seelim_stub.xlsx"
    _write_named_stub(named)
    if not is_stub_workbook(named):
        raise AssertionError("SeeLim '(stub)' A1 sheets must count as a stub")
    mixed = tmp / "mixed_stub.xlsx"
    wb_m = Workbook()
    ws_m = wb_m.active
    ws_m.title = "Summary"
    ws_m["A1"] = "Summary (stub)"
    leak = wb_m.create_sheet("LeakageOff")
    leak["A1"] = "LeakageOff (stub)"
    ron = wb_m.create_sheet("Ron")
    ron["A1"] = "Parameter"
    wb_m.save(mixed)
    wb_m.close()
    if not is_stub_workbook(mixed):
        raise AssertionError("SeeLim (stub)+Parameter mix must stay a stub")

    import yaml
    from ate.core import golden_refs as gr

    dummy_g = tmp / "msop_golden.xlsx"
    _write_senior(dummy_g)
    idx = tmp / "index.yaml"
    idx.write_text(
        yaml.safe_dump(
            {
                "parts": {
                    "rs2323": {
                        "default": "MSOP",
                        "packages": {"MSOP": {"golden": str(dummy_g)}},
                    }
                }
            },
            sort_keys=False,
        ),
        encoding="utf-8",
    )
    old_idx = gr.index_path
    gr.index_path = lambda: idx
    try:
        if gr.resolve_golden("rs2323", "MSOP") is None:
            raise AssertionError("MSOP golden must resolve")
        if gr.resolve_golden("rs2323", "UQFN1.4X1.8-10") is not None:
            raise AssertionError("UQFN must not fall back to MSOP golden")
        if gr.resolve_golden("rs2323") is None:
            raise AssertionError("empty package still uses default MSOP")
    finally:
        gr.index_path = old_idx

    info = inspect_workbook(senior)
    if info.get("stub"):
        raise AssertionError("VIX/VOX Test Parameter book must not be stub")
    if "VIX" not in (info.get("sheets") or []) or "VOX" not in (info.get("sheets") or []):
        raise AssertionError(f"senior sheets missing VIX/VOX: {info.get('sheets')}")

    rows = [
        {
            "part_key": "rs1g08",
            "package": "SC70-5",
            "stub": True,
            "score": 0,
            "path": str(stub),
            "sheets": ["VIH_VIL", "VOH"],
        },
        {
            "part_key": "rs1g08",
            "package": "SOT23",
            "stub": False,
            "score": 5000,
            "path": str(senior),
            "sheets": ["Summary", "VIX", "VOX"],
        },
    ]
    best = _best_by_key(rows)
    hit = best.get("rs1g08") or best.get("rs1g08|SOT23")
    if hit is None or "VIX" not in (hit.get("sheets") or []):
        raise AssertionError(f"best must pick SOT23 VIX book, got {best}")
    if str(hit.get("path")) == str(stub):
        raise AssertionError("JianHong stub must not win")

    live = Path(TEST_DB_ROOT) / "Logic" / "RS1G08"
    if live.is_dir():
        found_senior = False
        for xlsx in live.rglob("*Lab_Report.xlsx"):
            if "_filled" in xlsx.name.lower() or "_ate" in xlsx.parts:
                continue
            got = inspect_workbook(xlsx)
            if got.get("stub"):
                continue
            names = {_fold_sheet(n) for n in (got.get("sheets") or [])}
            if "vix" in names and "vox" in names:
                found_senior = True
                break
        if not found_senior:
            raise AssertionError("live RS1G08 has no senior VIX/VOX Lab_Report")

    golden = resolve_golden("rs1g08", "SOT23") or resolve_golden("rs1g08")
    if golden is not None:
        ginfo = inspect_workbook(golden)
        names = {_fold_sheet(n) for n in (ginfo.get("sheets") or [])}
        if ginfo.get("stub"):
            raise AssertionError(f"stored golden is a stub: {golden}")
        if "vix" not in names or "vox" not in names:
            raise AssertionError(f"stored RS1G08 golden missing VIX/VOX: {ginfo.get('sheets')}")

    from ate.core.lookup import reference_root
    from ate.core.new_product import reports_root

    drop = reports_root()
    drop_s = str(drop)
    if "Product Testing Report" not in drop_s and drop.name != "goldens":
        raise AssertionError(f"reports_root must be Product Testing Report or _ate/goldens, got {drop}")
    ref = reference_root()
    if ref.name != "Reference" or ref.parent.name != "Reference":
        raise AssertionError(f"datasheets must be Downloads/Reference/Reference, got {ref}")
    for stem in ("RS1GT08_(RevA.2.1).pdf", "RS1GT32_(RevA.2.1).pdf"):
        if not (ref / stem).is_file():
            raise AssertionError(f"datasheet missing: {ref / stem}")
    for pk in ("rs1gt08", "rs1gt32"):
        g = resolve_golden(pk, "SC70-5") or resolve_golden(pk)
        if g is None:
            raise AssertionError(f"golden missing for {pk} (need Product Testing Report + collect)")
        gi = inspect_workbook(g)
        names = {_fold_sheet(n) for n in (gi.get("sheets") or [])}
        if gi.get("stub") or "vix" not in names or "vox" not in names:
            raise AssertionError(f"{pk} golden must be VIX/VOX, got {gi.get('sheets')}")
    for part in ("RS1GT08", "RS1GT32"):
        book = Path(TEST_DB_ROOT) / "Logic" / part / "SC70-5" / "Ariff" / "Version_1" / "workbook" / f"{part}_Lab_Report.xlsx"
        if not book.is_file():
            raise AssertionError(f"Ariff workbook missing: {book}")
        ai = inspect_workbook(book)
        names = {_fold_sheet(n) for n in (ai.get("sheets") or [])}
        if ai.get("stub") or "vix" not in names:
            raise AssertionError(f"Ariff {part} must be VIX/VOX lab book")

    print("OK check_golden_refs stub=skip rs1g08=VIX/VOX rs1gt08/32=Ariff")
    return 0


def _fold_sheet(name: str) -> str:
    return "".join(ch for ch in str(name or "").lower() if ch.isalnum())


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"FAIL check_golden_refs: {exc}", file=sys.stderr)
        raise SystemExit(1)
