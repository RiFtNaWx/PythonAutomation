"""Fail-closed: datasheet ingest (OCR only if thin), golden Excel copy, sweep sort.

Run: python -m ate.core.check_ingest_datasheet
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path


def main() -> int:
    errors: list[str] = []
    from ate.core.ocr_engine import extract_pdf, text_is_thin

    if not text_is_thin("x" * 50):
        errors.append("short blob must be thin")
    if text_is_thin("VOH " * 80):
        errors.append("fat datasheet text must not force OCR")
    try:
        extract_pdf(Path("missing.pdf"), engine="qianfan")
        errors.append("qianfan must refuse before download")
    except (FileNotFoundError, ValueError):
        pass
    try:
        extract_pdf(Path(__file__), engine="qianfan", named_cloud=True)
        errors.append("named qianfan must still refuse 5B download")
    except (ValueError, FileNotFoundError):
        pass

    tmp = Path(tempfile.mkdtemp(prefix="ate_ingest_"))
    from openpyxl import Workbook, load_workbook

    xlsx = tmp / "lab.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "IDD"
    ws["D10"] = None
    wb.save(xlsx)
    wb.close()

    class _Ctx:
        def lab_report_path(self):
            return xlsx

        def load_sheet_map(self):
            return {
                "tests": {
                    "supply_current": {
                        "excel_sheet": "IDD",
                        "paste": {"values": {"ICC_uA": "D10"}},
                    }
                }
            }

        def sessions_dir(self):
            return tmp

        def sheet_map_path(self):
            return tmp / "sheet_map.yaml"

    report = {
        "steps": [
            {
                "test_id": "cin",
                "dut": 1,
                "channel": "CHA",
                "measurements": [{"id": "CIN_pF", "value": 4.0, "unit": "pF"}],
                "data": {
                    "rows": [
                        {"VCC": 3.3, "freq_hz": 10e6, "CIN_pF": 3.9},
                        {"VCC": 3.3, "freq_hz": 1e6, "CIN_pF": 4.1},
                        {"VCC": 3.3, "freq_hz": 5e6, "CIN_pF": 4.0},
                    ]
                },
            },
            {
                "test_id": "supply_current",
                "dut": 1,
                "measurements": [{"id": "ICC_uA", "value": 0.8}],
            },
        ]
    }
    from ate.core.ingest_datasheet import copilot_prompt, excel_map, organize_sweeps
    from ate.reporting.session_values import fill_workbook_from_report

    sweeps = organize_sweeps(report, ctx=_Ctx())
    if sweeps.get("blocks") != 1:
        errors.append(f"organize_sweeps blocks={sweeps}")
    else:
        freqs = [r.get("freq_hz") for r in sweeps["sweeps"][0]["rows"]]
        if freqs != [1e6, 5e6, 10e6]:
            errors.append(f"sweep sort want 1/5/10 MHz, got {freqs}")
    md = Path(sweeps["markdown"]).read_text(encoding="utf-8")
    if "CIN_pF" not in md or "cin" not in md:
        errors.append("sweep_summary.md missing cin table")

    mapped = excel_map(ctx=_Ctx())
    ids = [m["id"] for m in mapped.get("mapped") or []]
    if "ICC_uA" not in ids:
        errors.append(f"excel_map missed ICC_uA {mapped}")

    res = fill_workbook_from_report(
        report=report, workbook_path=xlsx, ctx=_Ctx(), copy_golden=True
    )
    filled = xlsx.with_name("lab_filled.xlsx")
    if res.get("status") == "error":
        errors.append(f"copy_golden fill error {res}")
    elif not filled.is_file():
        errors.append("copy_golden must write *_filled.xlsx")
    else:
        live = load_workbook(xlsx, data_only=False)
        if live["IDD"]["D10"].value not in (None,):
            errors.append(f"golden source mutated D10={live['IDD']['D10'].value!r}")
        live.close()
        copy = load_workbook(filled, data_only=False)
        if copy["IDD"]["D10"].value != 0.8:
            errors.append(f"filled D10={copy['IDD']['D10'].value!r} want 0.8")
        if "Sweep" not in copy.sheetnames:
            errors.append("filled copy must gain Sweep sheet")
        else:
            # row1 headers, row2 should be 1 MHz after sort
            freq_col = None
            for col in range(1, 12):
                if str(copy["Sweep"].cell(1, col).value or "") == "freq_hz":
                    freq_col = col
                    break
            if freq_col is None:
                errors.append("Sweep sheet missing freq_hz header")
            elif copy["Sweep"].cell(2, freq_col).value != 1e6:
                errors.append(
                    f"Sweep first data row freq={copy['Sweep'].cell(2, freq_col).value!r} want 1e6"
                )
        copy.close()

    prompt = copilot_prompt(
        part="RS1G07",
        limits={"specs": [{"id": "CIN_pF"}, {"id": "ICC_uA"}]},
        excel=mapped,
        extract={"engine": "text", "text_len": 900, "thin": False},
        fill=res,
        sweeps=sweeps,
    )
    if "Is it like this?" in prompt:
        errors.append("ingest must not ask to paste a Copilot block")
    if "AGENTS.md" not in prompt or "CIN_pF" not in prompt:
        errors.append("ingest log must name AGENTS.md and list spec ids")
    if "ICC_uA" in ids and "Unmapped: (none)" not in prompt and "ICC_uA" in prompt:
        pass
    if "Do not guess A91" not in prompt:
        errors.append("copilot must forbid guessed Excel cells")
    if "Slice at #Test_Database" not in prompt:
        errors.append("ingest prompt must include PATH_RULE so later AI rebases xlsx")
    from ate.core.ingest_datasheet import TABLES_DIR, export_part_tables
    from ate.core.paths import PATH_RULE as _PR

    if "ate/config/limits/" not in _PR or "truth_table" not in _PR:
        errors.append("PATH_RULE must name limits yaml and parts truth_table as SoT")
    idx = export_part_tables("RS1G08", xlsx_rel="RS1G08/RS1G08_Lab_Report.xlsx")
    if not (TABLES_DIR / "rs1g08.json").is_file():
        errors.append("export_part_tables must write datasheets/tables/<key>.json")
    elif not (idx.get("truth_table") or {}).get("rows"):
        errors.append("rs1g08 tables json must copy truth_table rows from parts yaml")
    if "C:\\Users\\" in str(idx.get("canonical") or {}):
        errors.append("tables json canonical paths must stay repo-relative")

    from ate.core.ingest_datasheet import ingest
    src = Path(__file__).resolve().parent / "ingest_datasheet.py"
    ingest_src = src.read_text(encoding="utf-8")
    if "--audit" not in ingest_src:
        errors.append("ingest_datasheet CLI must support --audit")
    if "--xlsx" not in ingest_src:
        errors.append("ingest_datasheet CLI must support --xlsx")
    from ate.core.lookup import resolve_pdf

    repo = Path(__file__).resolve().parents[2]
    if not (repo / "CLAUDE.md").is_file() or not (repo / "AI.md").is_file():
        errors.append("CLAUDE.md and AI.md must exist as the always-on system prompt")
    agents = (repo / "AGENTS.md").read_text(encoding="utf-8")
    if "Do not ask anyone to paste a Copilot block" not in agents:
        errors.append("AGENTS.md must forbid paste-copilot")

    downloads_xlsx = Path.home() / "Downloads" / "RS0204 Standard_Lab_Report.xlsx"
    downloads_pdf = Path.home() / "Downloads" / "RS0204_(RevA.5).pdf"
    if not downloads_xlsx.is_file():
        errors.append(f"RS0204 lab xlsx missing: {downloads_xlsx}")
    else:
        from ate.core.campaign_outline import _probe_a1_cell

        wb2 = load_workbook(downloads_xlsx, data_only=False)
        try:
            if _probe_a1_cell(wb2["VIH"]) != "K16":
                errors.append(f"VIH DUT1 A1 must probe K16, got {_probe_a1_cell(wb2['VIH'])}")
            if _probe_a1_cell(wb2["VIL"]) != "K15":
                errors.append(f"VIL DUT1 A1 must probe K15, got {_probe_a1_cell(wb2['VIL'])}")
            if _probe_a1_cell(wb2["Il"]) != "G12":
                errors.append(f"Il DUT1 A1 must probe G12, got {_probe_a1_cell(wb2['Il'])}")
        finally:
            wb2.close()
    if not downloads_pdf.is_file():
        errors.append(f"RS0204 PDF missing: {downloads_pdf}")
    else:
        got_pdf = extract_pdf(downloads_pdf, engine="text")
        blob = str(got_pdf.get("text") or "")
        if text_is_thin(blob):
            errors.append("RS0204 RevA.5 PDF text is thin; OCR loop would fire (want local text)")
        if "0204" not in blob:
            errors.append("RS0204 PDF text extract missed part number")

    pdf = resolve_pdf("RS1G07")
    if pdf is not None and Path(pdf).is_file():
        got = ingest(
            "RS1G07",
            web_ok=False,
            ocr="off",
            fill_excel=True,
            copy_golden=True,
            ctx=_Ctx(),
        )
        eng = (got.get("extract") or {}).get("engine")
        if eng not in ("text", "none"):
            errors.append(f"RS1G07 with text PDF must not OCR, engine={eng}")
        specs = (got.get("limits") or {}).get("specs") or []
        if not any(str(s.get("id")) == "ICC_uA" for s in specs if isinstance(s, dict)):
            errors.append("ingest RS1G07 must keep ICC_uA in limits")
        if "Is it like this?" in str(got.get("copilot") or ""):
            errors.append("ingest must not ask to paste Copilot")
        if "AGENTS.md" not in str(got.get("copilot") or ""):
            errors.append("ingest log must name AGENTS.md")
    else:
        errors.append("RS1G07 PDF missing from local Reference (cannot prove ingest)")

    if downloads_xlsx.is_file() and downloads_pdf.is_file():
        got0204 = ingest(
            "RS0204",
            pdf=downloads_pdf,
            xlsx=downloads_xlsx,
            web_ok=False,
            ocr="off",
            fill_excel=False,
            copy_golden=True,
        )
        eng0204 = (got0204.get("extract") or {}).get("engine")
        if eng0204 not in ("text", "none"):
            errors.append(f"RS0204 PDF must stay on local text, engine={eng0204}")
        specs0204 = (got0204.get("limits") or {}).get("specs") or []
        if not any(str(s.get("id")) == "ICC_uA" and s.get("max") == 10 for s in specs0204 if isinstance(s, dict)):
            errors.append("ingest RS0204 must keep ICC_uA max 10")
        mapped0204 = {
            str(m.get("id")): m.get("cells")
            for m in (got0204.get("excel") or {}).get("mapped") or []
            if isinstance(m, dict)
        }
        for mid, cell in (("VIH_V", "K16"), ("VIL_V", "K15"), ("IL_uA", "G12")):
            if mapped0204.get(mid) != cell:
                errors.append(f"RS0204 sheet_map {mid} want {cell} got {mapped0204.get(mid)}")

    if errors:
        print("FAIL check_ingest_datasheet:")
        for e in errors:
            print(f"  - {e}")
        return 1
    print("OK check_ingest_datasheet: text-first OCR, golden *_filled.xlsx, RS0204 K16/K15/G12, AGENTS.md")
    return 0


if __name__ == "__main__":
    sys.exit(main())
