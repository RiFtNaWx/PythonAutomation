"""OpAmp lab workbook vs sheet_map vs TestSpec.lab_sheet sync (A05-T01).

Run: python -m ate.core.check_lab_report_sync
"""
from __future__ import annotations

import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook

from ate.core.database import default_context
from ate.core.registry import all_tests, load_family

_NON_TEST_SHEETS = frozenset({"Summary", "Checklist"})


def _workbook_test_sheets(xlsx_path: Path) -> set[str]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        return {s for s in wb.sheetnames if s not in _NON_TEST_SHEETS}
    finally:
        wb.close()


def _map_excel_sheets(sheet_map: dict) -> dict[str, str]:
    """test_key -> excel_sheet from sheet_map.tests."""
    tests = sheet_map.get("tests") or {}
    out: dict[str, str] = {}
    if not isinstance(tests, dict):
        return out
    for key, entry in tests.items():
        if not isinstance(entry, dict):
            continue
        sheet = entry.get("excel_sheet")
        if sheet:
            out[str(key)] = str(sheet)
    return out


def check_lab_report_sync() -> list[str]:
    """Return human-readable failure lines (empty = pass)."""
    load_family("opamp")
    registered = {t.lab_sheet for t in all_tests() if t.lab_sheet}

    ctx = default_context()
    xlsx = ctx.lab_report_path()
    if not xlsx.is_file():
        return [f"workbook missing: {xlsx}"]

    sm_path = ctx.sheet_map_path()
    if not sm_path.is_file():
        return [f"sheet_map missing: {sm_path}"]

    with sm_path.open(encoding="utf-8") as fh:
        sheet_map = yaml.safe_load(fh) or {}

    wb_sheets = _workbook_test_sheets(xlsx)
    map_sheets = _map_excel_sheets(sheet_map)
    errors: list[str] = []

    # (1) workbook test sheet with no registered lab_sheet
    unregistered = sorted(wb_sheets - registered)
    if unregistered:
        errors.append(
            f"workbook sheets without TestSpec.lab_sheet: {unregistered}"
        )

    # (2) registered lab_sheet missing from workbook
    missing_reg = sorted(registered - wb_sheets)
    if missing_reg:
        errors.append(
            f"registered lab_sheet missing from workbook: {missing_reg}"
        )

    # (3) sheet_map excel_sheet not on workbook (drift / orphan claim)
    bad_map: list[str] = []
    for test_key, excel_sheet in sorted(map_sheets.items()):
        if excel_sheet not in wb_sheets:
            bad_map.append(f"{test_key} -> {excel_sheet!r}")
    if bad_map:
        errors.append(
            "sheet_map excel_sheet not on workbook: " + "; ".join(bad_map)
        )

    # (4) sheet_map excel_sheet with no TestSpec.lab_sheet
    mapped_unreg = sorted(set(map_sheets.values()) - registered)
    if mapped_unreg:
        errors.append(
            "sheet_map excel_sheet without TestSpec.lab_sheet: " + ", ".join(mapped_unreg)
        )

    # (5) mapped lab sheets must not still be RuntimeError stubs
    stub_files = []
    for spec in all_tests():
        if not spec.lab_sheet or spec.lab_sheet not in set(map_sheets.values()):
            continue
        src = getattr(spec.run, "__code__", None)
        filename = getattr(src, "co_filename", "") if src else ""
        if filename.replace("\\", "/").endswith("/stubs.py"):
            stub_files.append(spec.id)
    if stub_files:
        errors.append("mapped tests still registered from stubs.py: " + ", ".join(stub_files))

    return errors


def main() -> int:
    errors = check_lab_report_sync()
    if errors:
        print("FAIL lab-report sync:")
        for line in errors:
            print(f"  - {line}")
        return 1
    print("OK lab-report sync: workbook, sheet_map, and OpAmp lab_sheet agree")
    return 0


if __name__ == "__main__":
    sys.exit(main())
