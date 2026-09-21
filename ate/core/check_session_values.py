"""Self-check: sheet_map paste.values fills workbook cells from report.json.

Run: python -m ate.core.check_session_values
"""
from __future__ import annotations

import sys
import tempfile
from pathlib import Path


def main() -> int:
    errors: list[str] = []
    tmp = Path(tempfile.mkdtemp(prefix="ate_a19_values_"))
    from openpyxl import Workbook, load_workbook

    xlsx = tmp / "lab.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "GBW"
    ws["B12"] = None
    ws["B13"] = None
    wb.save(xlsx)
    wb.close()

    class _Ctx:
        def lab_report_path(self):
            return xlsx

        def load_sheet_map(self):
            return {
                "tests": {
                    "GBW": {
                        "excel_sheet": "GBW",
                        "folder": "GBW",
                        "paste": {"values": {"GBW_MHz": "B12", "result": "B13"}},
                    }
                }
            }

    report = {
        "steps": [
            {
                "test_id": "gbw",
                "dut": 1,
                "success": True,
                "measurements": [{"id": "GBW_MHz", "value": 7.0, "result": "unspec"}],
            }
        ]
    }
    from ate.reporting.session_values import fill_workbook_from_report

    res = fill_workbook_from_report(report=report, workbook_path=xlsx, ctx=_Ctx())
    if res.get("filled", 0) < 1:
        errors.append(f"expected fill, got {res}")
    else:
        wb2 = load_workbook(xlsx, data_only=False)
        if wb2["GBW"]["B12"].value != 7.0:
            errors.append(f"B12={wb2['GBW']['B12'].value!r} want 7.0")
        if str(wb2["GBW"]["B13"].value or "").upper() != "UNSPEC":
            errors.append(f"B13={wb2['GBW']['B13'].value!r} want UNSPEC")
        wb2.close()

    class _DutCtx:
        def lab_report_path(self):
            return xlsx

        def load_sheet_map(self):
            return {
                "tests": {
                    "GBW": {
                        "excel_sheet": "GBW",
                        "folder": "GBW",
                        "paste": {
                            "values": {
                                "GBW_MHz": {
                                    "CHA": ["B12", "C12", "D12", "E12"],
                                    "CHB": ["B14", "C14", "D14", "E14"],
                                }
                            }
                        },
                    }
                }
            }

    wb3 = load_workbook(xlsx)
    wb3["GBW"]["C12"] = None
    wb3["GBW"]["B14"] = None
    wb3.save(xlsx)
    wb3.close()
    res2 = fill_workbook_from_report(
        report={
            "steps": [
                {
                    "test_id": "gbw",
                    "dut": 2,
                    "channel": "CHA",
                    "measurements": [{"id": "GBW_MHz", "value": 6.5}],
                },
                {
                    "test_id": "gbw",
                    "dut": 1,
                    "channel": "CHB",
                    "measurements": [{"id": "GBW_MHz", "value": 8.25}],
                },
            ]
        },
        workbook_path=xlsx,
        ctx=_DutCtx(),
    )
    if res2.get("filled", 0) < 2:
        errors.append(f"DUT/channel fill expected 2 writes, got {res2}")
    else:
        wb4 = load_workbook(xlsx, data_only=False)
        if wb4["GBW"]["C12"].value != 6.5:
            errors.append(f"C12 DUT2 CHA={wb4['GBW']['C12'].value!r} want 6.5")
        if wb4["GBW"]["B14"].value != 8.25:
            errors.append(f"B14 DUT1 CHB={wb4['GBW']['B14'].value!r} want 8.25")
        wb4.close()

    from ate.reporting.session_paste import _is_real_cell

    if _is_real_cell("FILL_ME"):
        errors.append("FILL_ME must skip")

    class _LogicCtx:
        def lab_report_path(self):
            return xlsx

        def load_sheet_map(self):
            return {
                "tests": {
                    "VOH": {
                        "excel_sheet": "VOH",
                        "folder": "VOH",
                        "paste": {"values": {"VOH_2p0V": ["B10", "C10", "D10", "E10"]}},
                    }
                }
            }

    wb5 = load_workbook(xlsx)
    if "VOH" not in wb5.sheetnames:
        ws = wb5.create_sheet("VOH")
    else:
        ws = wb5["VOH"]
    ws["B10"] = None
    wb5.save(xlsx)
    wb5.close()
    res3 = fill_workbook_from_report(
        report={
            "steps": [
                {
                    "test_id": "voh_load",
                    "dut": 1,
                    "measurements": [{"id": "VOH_2p0V", "value": 1.7}],
                }
            ]
        },
        workbook_path=xlsx,
        ctx=_LogicCtx(),
    )
    if res3.get("filled", 0) < 1:
        errors.append(f"voh_load must map to sheet VOH, got {res3}")
    else:
        wb6 = load_workbook(xlsx, data_only=False)
        if wb6["VOH"]["B10"].value != 1.7:
            errors.append(f"VOH!B10={wb6['VOH']['B10'].value!r} want 1.7")
        wb6.close()

    class _TrackCtx:
        def lab_report_path(self):
            return xlsx

        def load_sheet_map(self):
            return {
                "tests": {
                    "VOH": {
                        "excel_sheet": "VOX",
                        "folder": "VOH",
                        "paste": {"values": {"VOH_4p5V": ["G16", "H16", "I16"]}},
                    },
                    "Supply_Current": {
                        "excel_sheet": "ICC",
                        "folder": "Supply_Current",
                        "paste": {"values": {"ICC_uA": "D10"}},
                    },
                    "Iplus": {
                        "excel_sheet": "Iplus",
                        "folder": "Iplus",
                        "paste": {"values": {"IPLUS_uA": "B2"}},
                    },
                }
            }

    wb7 = load_workbook(xlsx)
    for name in ("VOX", "ICC", "Iplus"):
        if name not in wb7.sheetnames:
            wb7.create_sheet(name)
    wb7["VOX"]["G16"] = None
    wb7["ICC"]["D10"] = None
    wb7["Iplus"]["B2"] = None
    wb7.save(xlsx)
    wb7.close()
    res4 = fill_workbook_from_report(
        report={
            "steps": [
                {
                    "test_id": "voh_load",
                    "dut": 1,
                    "measurements": [{"id": "VOH_4p5V", "value": 4.2}],
                },
                {
                    "test_id": "supply_current_sweep",
                    "dut": 1,
                    "measurements": [{"id": "ICC_uA", "value": 0.9}],
                },
                {
                    "test_id": "iplus",
                    "dut": 1,
                    "measurements": [{"id": "IPLUS_uA", "value": 0.4}],
                },
            ]
        },
        workbook_path=xlsx,
        ctx=_TrackCtx(),
    )
    if res4.get("filled", 0) < 3:
        errors.append(f"VOX/ICC/Iplus fill expected 3 writes, got {res4}")
    else:
        wb8 = load_workbook(xlsx, data_only=False)
        if wb8["VOX"]["G16"].value != 4.2:
            errors.append(f"VOX!G16={wb8['VOX']['G16'].value!r} want 4.2")
        if wb8["ICC"]["D10"].value != 0.9:
            errors.append(f"ICC!D10={wb8['ICC']['D10'].value!r} want 0.9")
        if wb8["Iplus"]["B2"].value != 0.4:
            errors.append(f"Iplus!B2={wb8['Iplus']['B2'].value!r} want 0.4")
        wb8.close()

    wb9 = load_workbook(xlsx)
    if "GBW" not in wb9.sheetnames:
        wb9.create_sheet("GBW")
    wb9["GBW"].merge_cells("B12:C12")
    wb9["GBW"]["B12"] = None
    wb9.save(xlsx)
    wb9.close()

    class _MergeCtx:
        def lab_report_path(self):
            return xlsx

        def load_sheet_map(self):
            return {
                "tests": {
                    "GBW": {
                        "excel_sheet": "GBW",
                        "folder": "GBW",
                        "paste": {"values": {"GBW_MHz": "C12"}},
                    }
                }
            }

    res5 = fill_workbook_from_report(
        report={
            "steps": [
                {
                    "test_id": "gbw",
                    "dut": 1,
                    "measurements": [{"id": "GBW_MHz", "value": 7.1}],
                }
            ]
        },
        workbook_path=xlsx,
        ctx=_MergeCtx(),
    )
    if res5.get("status") == "error":
        errors.append(f"merged cell must skip not crash, got {res5}")

    from ate.core.datalog import _step_key

    if _step_key({"test_id": "vos_sweep", "dut": 1}) != _step_key(
        {"test_id": "vos_sweep", "dut": 1, "channel": "CHA"}
    ):
        errors.append("DEMO/live vos must merge: empty channel is CHA")

    wb10 = load_workbook(xlsx)
    if "IDD" not in wb10.sheetnames:
        wb10.create_sheet("IDD")
    wb10["IDD"]["D10"] = None
    wb10.save(xlsx)
    wb10.close()

    class _IddCtx:
        def lab_report_path(self):
            return xlsx

        def load_sheet_map(self):
            return {
                "tests": {
                    "IDD": {
                        "excel_sheet": "IDD",
                        "folder": "IDD",
                        "paste": {"values": {"ICC_uA": "D10"}},
                    }
                }
            }

    res6 = fill_workbook_from_report(
        report={
            "steps": [
                {
                    "test_id": "supply_current",
                    "dut": 1,
                    "measurements": [{"id": "ICC_uA", "value": 0.8}],
                }
            ]
        },
        workbook_path=xlsx,
        ctx=_IddCtx(),
    )
    if res6.get("filled", 0) < 1:
        errors.append(f"IDD sheet must match supply_current, got {res6}")
    else:
        wb11 = load_workbook(xlsx, data_only=False)
        if wb11["IDD"]["D10"].value != 0.8:
            errors.append(f"IDD!D10={wb11['IDD']['D10'].value!r} want 0.8")
        wb11.close()

    res7 = fill_workbook_from_report(
        report={
            "steps": [
                {
                    "test_id": "supply_current",
                    "dut": 1,
                    "measurements": [{"id": "ICC_uA", "value": 1.2}],
                }
            ]
        },
        workbook_path=xlsx,
        ctx=_IddCtx(),
        demo=True,
    )
    sidecar = xlsx.with_name(xlsx.stem + "_demo.xlsx")
    if res7.get("status") == "error":
        errors.append(f"demo fill error: {res7}")
    elif not sidecar.is_file():
        errors.append("demo fill must write *_demo.xlsx")
    else:
        wb_live = load_workbook(xlsx, data_only=False)
        if wb_live["IDD"]["D10"].value != 0.8:
            errors.append(
                f"demo fill must not change live xlsx, D10={wb_live['IDD']['D10'].value!r}"
            )
        wb_live.close()
        wb_demo = load_workbook(sidecar, data_only=False)
        if wb_demo["IDD"]["D10"].value != 1.2:
            errors.append(f"demo xlsx D10={wb_demo['IDD']['D10'].value!r} want 1.2")
        wb_demo.close()

    class _NarCtx:
        part_key = "rs622"

        def lab_report_path(self):
            return xlsx

        def load_sheet_map(self):
            return {
                "tests": {
                    "slew": {
                        "excel_sheet": "Slew Rate",
                        "folder": "Slew Rate",
                        "paste": {
                            "values": {"SR_Vus": "D20"},
                            "narrative": {
                                "conditions": "B3",
                                "datasheet": "B8",
                                "conclusion": "B15",
                                "circuitry": "B4",
                            },
                        },
                    }
                }
            }

    wb_n = load_workbook(xlsx)
    if "Slew Rate" not in wb_n.sheetnames:
        ws_n = wb_n.create_sheet("Slew Rate")
    else:
        ws_n = wb_n["Slew Rate"]
    ws_n["A3"] = "Test Conditions"
    ws_n["B3"] = "#VALUE!"
    ws_n["A4"] = "Test Circuitry"
    ws_n["B4"] = "#VALUE!"
    ws_n["A8"] = "Datasheet"
    ws_n["B8"] = "#VALUE!"
    ws_n["A15"] = "Test Conclusion"
    ws_n["B15"] = None
    ws_n["D20"] = None
    wb_n.save(xlsx)
    wb_n.close()
    res8 = fill_workbook_from_report(
        report={
            "steps": [
                {
                    "test_id": "slew",
                    "dut": 1,
                    "channel": "CHA",
                    "measurements": [{"id": "SR_Vus", "value": 4.05, "unit": "V/us"}],
                }
            ]
        },
        workbook_path=xlsx,
        ctx=_NarCtx(),
        copy_golden=True,
    )
    filled_path = xlsx.with_name(xlsx.stem + "_filled.xlsx")
    if res8.get("status") == "error":
        errors.append(f"narrative fill error: {res8}")
    elif not filled_path.is_file():
        errors.append("copy_golden narrative fill must write *_filled.xlsx")
    else:
        live_n = load_workbook(xlsx, data_only=False)
        if str(live_n["Slew Rate"]["B8"].value or "").strip().upper() != "#VALUE!":
            errors.append(
                f"narrative fill must not mutate golden source, B8={live_n['Slew Rate']['B8'].value!r}"
            )
        live_n.close()
        filled_n = load_workbook(filled_path, data_only=False)
        ds = str(filled_n["Slew Rate"]["B8"].value or "")
        if "SR_Vus" not in ds or "3.7" not in ds:
            errors.append(f"filled datasheet must crawl SR_Vus typ 3.7, got {ds!r}")
        conc = str(filled_n["Slew Rate"]["B15"].value or "")
        if "4.05" not in conc or "SR_Vus" not in conc:
            errors.append(f"filled conclusion must use measured 4.05, got {conc!r}")
        if str(filled_n["Slew Rate"]["B4"].value or "").strip().upper() == "#VALUE!":
            errors.append("filled circuitry must clear #VALUE!")
        cond = str(filled_n["Slew Rate"]["B3"].value or "")
        if "BUFFER" not in cond.upper() and "buffer" not in cond.lower() and "10" not in cond:
            errors.append(f"filled conditions must use BUFFER checklist, got {cond!r}")
        if filled_n["Slew Rate"]["D20"].value != 4.05:
            errors.append(f"filled D20={filled_n['Slew Rate']['D20'].value!r} want 4.05")
        if filled_n["Slew Rate"]["D20"].comment is None:
            errors.append("filled number cell must have an ATE comment")
        filled_n.close()

    if errors:
        print("FAIL check_session_values:")
        for e in errors:
            print(f"  - {e}")
        return 1
    print("OK check_session_values")
    return 0


if __name__ == "__main__":
    sys.exit(main())
