"""Map Product Testing Report sheets -> existing TestSpecs. Local files only.

Does not scrape en.run-ic.com into #Test_Database. Does not guess Excel cells.
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

import yaml

from ate.core.lookup import TEXT_DIR, resolve_pdf
from ate.core.paths import CONFIG_DIR, PATH_RULE, TEST_DB_ROOT, resolve_portable, store_portable, store_rel
from ate.core.specs import LIMITS_DIR, load_part_yaml

COVERAGE_PATH = CONFIG_DIR / "datasheets" / "coverage.json"

_SKIP_FOLD = frozenset(
    {
        "summary",
        "checklist",
        "status",
        "problems",
        "question",
        "sheet1",
        "sheet2",
        "mapping",
        "cover",
        "info",
        "datalogrd",
        "iccrd",
        "iscrd",
        "aolrd",
        "psrrrd",
        "cmrrrd",
        "gbwretest",
        "chippick",
    }
)
# OpAmp pages copied into some Logic lab books -- not Logic TestSpecs.
_LOGIC_LEFTOVER = frozenset({"sr", "psrr", "settingtime", "slewrate", "trtf"})

# Folded sheet name -> TestSpec ids that already exist in this repo.
SHEET_TO_IDS: dict[str, tuple[str, ...]] = {
    "vix": ("vih_vil",),
    "vihvil": ("vih_vil",),
    "vihl": ("vih_vil",),
    "vih": ("vih_vil", "vih"),
    "vil": ("vih_vil", "vil"),
    "vox": ("voh_load", "vol_load"),
    "voh": ("voh_load", "voh"),
    "vol": ("vol_load", "vol", "output_voltage", "vohl"),
    "vohl": ("voh_load", "vol_load", "vohl"),
    "vth": ("input_thresholds", "vih_vil"),
    "dvt": ("vih_vil",),
    "icc": ("supply_current", "supply_current_sweep", "icc"),
    "ic": ("supply_current",),
    "idd": ("supply_current",),
    "icct": ("delta_supply_current",),
    "dicc": ("delta_supply_current",),
    "iil": ("input_leakage_sweep",),
    "il": ("input_leakage_sweep", "il", "input_leakage"),
    "ii": ("input_leakage_sweep",),
    "iin": ("input_leakage_sweep", "input_leakage"),
    "ioff": ("ioff_leakage",),
    "cin": ("cin",),
    "ci": ("cin",),
    "cpd": ("cpd",),
    "cio": ("cpd",),
    "tpd": ("tp", "tpd", "tp_rs0204"),
    "tp": ("tp",),
    "ten": ("ten",),
    "tdis": ("tdis",),
    "tidle": ("tidle",),
    "slewrate": ("slew",),
    "sr": ("slew",),
    "gbw": ("gbw",),
    "vos": ("vos_sweep",),
    "settlingtime": ("settling",),
    "sssr": ("sssr",),
    "lssr": ("lssr",),
    "ort": ("ort",),
    "psrr": ("psrr",),
    "cmrr": ("cmrr",),
    "aol": ("aol",),
    "emirr": ("emirr",),
    "nophasereversal": ("no_phase_reversal",),
    "powerontime": ("power_on_time",),
    "noise": ("noise",),
    "iq": ("iq",),
    "iplus": ("iplus",),
    "isc": ("isc",),
    "vprecharge": ("output_voltage",),
}


def _fold(raw: str) -> str:
    s = str(raw or "").strip().lower()
    s = s.replace("∆", "d").replace("Δ", "d").replace("δ", "d")
    s = s.replace("μ", "u").replace("µ", "u")
    return re.sub(r"[^a-z0-9]+", "", s)


def _skip_sheet(name: str, family: str) -> str:
    fold = _fold(name)
    if not fold or fold in _SKIP_FOLD:
        return "meta"
    if "gaowen" in fold or fold.endswith("rd") and fold not in ("ort",):
        return "retest"
    if family in ("logic", "switch", "power") and fold in _LOGIC_LEFTOVER:
        return "opamp_leftover"
    if "gaodiwen" in fold:
        return "retest"
    return ""


def _ids_for_sheet(name: str) -> tuple[str, ...]:
    return SHEET_TO_IDS.get(_fold(name), ())


def _enabled_ids(part_key: str) -> list[str]:
    data = load_part_yaml(part_key)
    raw = data.get("enabled_tests")
    if isinstance(raw, list) and raw:
        return [str(x) for x in raw]
    out: list[str] = []
    modes = data.get("fixture_modes") or {}
    if isinstance(modes, dict):
        for mode in modes.values():
            if not isinstance(mode, dict):
                continue
            tests = mode.get("tests") or []
            if isinstance(tests, list):
                out.extend(str(x) for x in tests)
    return out


def _spec_ids(part_key: str) -> list[str]:
    path = LIMITS_DIR / f"{part_key}.yaml"
    if not path.is_file():
        return []
    blob = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    if not isinstance(blob, dict):
        return []
    return [str(s.get("id")) for s in (blob.get("specs") or []) if isinstance(s, dict) and s.get("id")]


def _is_datalog(rel: str) -> bool:
    s = rel.lower().replace("/", "\\")
    if "datalog" in s:
        return True
    if "\\files\\" in s or "\\raw data\\" in s:
        return True
    if "qualification" in s:
        return True
    return False


def _lab_xlsx(root: Path, report: str, *, part: str = "", package: str = "") -> Path | None:
    raw = str(report or "").strip()
    roots = [root]
    goldens = Path(TEST_DB_ROOT) / "_ate" / "goldens"
    if goldens not in roots:
        roots.append(goldens)
    for base in roots:
        if not base.is_dir() or not raw:
            continue
        cand = base / raw
        if cand.is_file() and cand.suffix.lower() == ".xlsx":
            return cand
        if cand.is_dir():
            labs = [
                p
                for p in cand.rglob("*.xlsx")
                if p.is_file() and not p.name.startswith("~") and not _is_datalog(str(p.relative_to(base)))
            ]
            if not labs:
                continue
            labs.sort(key=lambda p: (0 if "lab_report" in p.name.lower() or "test report" in p.name.lower() else 1, len(p.name)))
            return labs[0]
    sku = str(part or "").strip()
    if sku:
        try:
            from ate.core.golden_refs import resolve_golden

            g = resolve_golden(sku, package) or resolve_golden(sku)
        except Exception:
            g = None
        if g is not None:
            hit = resolve_portable(str(g), goldens, db_root=Path(TEST_DB_ROOT))
            if hit is not None and hit.is_file():
                return hit
            if g.is_file():
                return g
        if goldens.is_dir():
            needle = sku.lower()
            named = [
                p
                for p in goldens.rglob("golden.xlsx")
                if needle in str(p).replace("\\", "/").lower()
            ]
            if named:
                named.sort(key=lambda p: len(p.parts))
                return named[0]
    return None


def _sheet_names(xlsx: Path) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=False)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _registry_ids(family: str) -> set[str]:
    from ate.core.registry import all_tests, load_family

    fam = family
    if fam == "analog_switch":
        fam = "switch"
    if fam == "level":
        fam = "logic"
    load_family(fam if fam in ("opamp", "logic", "switch", "power") else "logic")
    return {t.id for t in all_tests()}


def audit(*, write: bool = True) -> dict[str, Any]:
    from ate.core.ingest_datasheet import export_part_tables
    from ate.core.new_product import reports_root

    inv_path = CONFIG_DIR / "inventory.yaml"
    inv = yaml.safe_load(inv_path.read_text(encoding="utf-8")) or {}
    root = reports_root()
    seen: set[str] = set()
    rows: list[dict[str, Any]] = []
    formats: dict[str, int] = {}
    for rec in inv.get("parts") or []:
        if not isinstance(rec, dict):
            continue
        sku = str(rec.get("part") or "").strip().upper()
        if not sku or sku in seen:
            continue
        seen.add(sku)
        pk = sku.lower()
        cat = str(rec.get("category") or "")
        family = {"analog_switch": "switch", "level": "level", "power": "power", "opamp": "opamp"}.get(cat, "logic")
        xlsx = _lab_xlsx(
            root,
            str(rec.get("report") or ""),
            part=sku,
            package=str(rec.get("package") or ""),
        )
        sheets: list[str] = []
        err = ""
        if xlsx is not None:
            try:
                sheets = _sheet_names(xlsx)
            except Exception as exc:
                err = str(exc)
        fmt = "none"
        if sheets:
            names = {_fold(s) for s in sheets}
            if "vox" in names or "vix" in names:
                fmt = "logic_lab_vix_vox"
            elif "vih" in names and "voh" in names:
                fmt = "logic_split_vih_voh"
            elif "slewrate" in names or "gbw" in names:
                fmt = "opamp_lab"
            elif "ron" in names:
                fmt = "switch_ron"
            elif "iq" in names:
                fmt = "ldo"
            else:
                fmt = "other"
        formats[fmt] = formats.get(fmt, 0) + 1
        enabled = _enabled_ids(pk)
        txt = TEXT_DIR / f"{pk}.txt"
        extract_len = len(txt.read_text(encoding="utf-8")) if txt.is_file() else 0
        pdf = resolve_pdf(sku)
        mapped: list[dict[str, Any]] = []
        missing_code: list[str] = []
        missing_enable: list[str] = []
        leftover: list[str] = []
        reg = _registry_ids(family)
        for name in sheets:
            why = _skip_sheet(name, family)
            if why == "opamp_leftover":
                leftover.append(name)
                continue
            if why:
                continue
            ids = tuple(i for i in _ids_for_sheet(name) if i in reg)
            if not ids:
                missing_code.append(name)
                continue
            hit = [i for i in ids if i in enabled]
            mapped.append({"sheet": name, "ids": list(ids), "enabled": hit})
            if not hit:
                missing_enable.append(name)
        export_part_tables(sku, xlsx_rel=str(rec.get("report") or ""))
        rows.append(
            {
                "part": sku,
                "part_key": pk,
                "category": cat,
                "family": family,
                "format": fmt,
                "xlsx": str(rec.get("report") or "") or (store_rel(xlsx, root) if xlsx else ""),
                "sheets": sheets,
                "enabled": enabled,
                "spec_ids": _spec_ids(pk),
                "extract_len": extract_len,
                "pdf": bool(pdf and Path(pdf).is_file()),
                "thin": extract_len < 200,
                "mapped": mapped,
                "missing_enable": missing_enable,
                "missing_code": missing_code,
                "opamp_leftover": leftover,
                "error": err,
            }
        )
    out = {
        "path_rule": PATH_RULE,
        "reports_root": store_portable(root) if root else "%USERPROFILE%/Downloads/Product Testing Report",
        "formats": formats,
        "parts": rows,
        "n": len(rows),
    }
    if write:
        COVERAGE_PATH.parent.mkdir(parents=True, exist_ok=True)
        COVERAGE_PATH.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
        out["path"] = str(COVERAGE_PATH)
    return out
