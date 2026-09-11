"""Campaign tags: _manifest/tags.yaml SoT + TAGS.txt for grep/Explorer.

Not a new folder axis. Board vocabulary lives in ate/config/boards.yaml.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Optional

import yaml

from ate.core.paths import CONFIG_DIR, TEST_DB_ROOT

_KIND_RE = re.compile(r"^[A-Za-z0-9_]+$")
_TAG_RE = re.compile(r"^[A-Za-z0-9_.:\-/]+$")
BOARDS_PATH = CONFIG_DIR / "boards.yaml"
LABEL_VOCAB_PATH = CONFIG_DIR / "label_vocab.yaml"
_SKIP_FAM_KEYS = frozenset({"board_types", "kinds", "default"})
_LABEL_VOCAB_PREAMBLE = (
    "# Shared label history (this class / all products).\n"
    "# Not a #Test_Database folder axis. Campaign tokens stay in _manifest/tags.yaml.\n"
    "\n"
)


def _norm_tag(raw: str) -> str:
    t = str(raw or "").strip()
    if not t:
        raise ValueError("empty tag")
    if not _TAG_RE.match(t):
        raise ValueError(f"invalid tag {t!r} (use letters, digits, _.:-/)")
    return t


def _norm_kind(raw: str) -> str:
    k = str(raw or "").strip().lower().replace(" ", "_")
    if k in ("board type", "board-type"):
        k = "board_type"
    if not k:
        k = "tag"
    if not _KIND_RE.match(k):
        raise ValueError(f"invalid label kind {raw!r}")
    return k


def _token(kind: str, value: str) -> str:
    if kind in ("tag", ""):
        return value
    return f"{kind}:{value}"


def _normalize_labels(raw: Any) -> list[dict[str, str]]:
    out: list[dict[str, str]] = []
    seen: set[tuple[str, str]] = set()

    def add(kind: str, value: str) -> None:
        try:
            k = _norm_kind(kind)
            v = _norm_tag(value)
        except ValueError:
            return
        key = (k, v.lower())
        if key in seen:
            return
        seen.add(key)
        out.append({"kind": k, "value": v})

    if isinstance(raw, list):
        for item in raw:
            if isinstance(item, dict):
                add(str(item.get("kind") or "tag"), str(item.get("value") or ""))
            elif isinstance(item, str) and str(item).strip():
                s = str(item).strip()
                if ":" in s:
                    k, _, v = s.partition(":")
                    add(k, v)
                else:
                    add("tag", s)
    return out


def _labels_from_stored(data: dict[str, Any]) -> list[dict[str, str]]:
    labels = _normalize_labels(data.get("labels"))
    if not labels:
        blob = {
            "tags": data.get("tags") or [],
            "boards": data.get("boards") or [],
        }
        extra: list[Any] = []
        for b in blob["boards"]:
            extra.append({"kind": "board", "value": str(b)})
        for t in blob["tags"]:
            extra.append(str(t))
        labels = _normalize_labels(extra)
    return labels


def tags_yaml_path(ctx) -> Path:
    return ctx.manifest_dir() / "tags.yaml"


def tags_txt_path(ctx) -> Path:
    return ctx.root() / "TAGS.txt"


def load_tags(ctx=None) -> dict[str, Any]:
    from ate.core.database import get_context

    c = ctx or get_context()
    path = tags_yaml_path(c)
    if not path.is_file():
        return {"tags": [], "boards": [], "labels": []}
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    if not isinstance(data, dict):
        return {"tags": [], "boards": [], "labels": []}
    labels = _labels_from_stored(data)
    boards = [x["value"] for x in labels if x["kind"] == "board"]
    tags = [_token(x["kind"], x["value"]) for x in labels]
    return {"tags": tags, "boards": boards, "labels": labels}


def save_tags(
    tags: list[str] | None = None,
    boards: list[str] | None = None,
    labels: list | None = None,
    *,
    ctx=None,
    remember_scope: str = "campaign",
) -> dict[str, Any]:
    """Write tags.yaml + TAGS.txt. Returns saved payload."""
    from ate.core.database import get_context

    c = ctx or get_context()
    c.ensure_tree()
    cur = load_tags(c)
    if labels is not None:
        parsed = _normalize_labels(labels)
    else:
        tag_list = list(tags) if tags is not None else list(cur.get("tags") or [])
        board_list = list(boards) if boards is not None else list(cur.get("boards") or [])
        extra: list[Any] = list(tag_list)
        extra.extend({"kind": "board", "value": str(b)} for b in board_list)
        parsed = _normalize_labels(extra)
    clean_boards = [x["value"] for x in parsed if x["kind"] == "board"]
    clean_tags = [_token(x["kind"], x["value"]) for x in parsed]
    payload = {"tags": clean_tags, "boards": clean_boards, "labels": parsed}
    ypath = tags_yaml_path(c)
    ypath.parent.mkdir(parents=True, exist_ok=True)
    ypath.write_text(
        yaml.safe_dump(payload, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )
    tpath = tags_txt_path(c)
    tpath.write_text("\n".join(clean_tags) + ("\n" if clean_tags else ""), encoding="utf-8")
    excel = sync_tags_to_workbook(c, clean_tags)
    remembered = remember_labels(parsed, scope=remember_scope, ctx=c)
    return {
        **payload,
        "tags_yaml": str(ypath),
        "tags_txt": str(tpath),
        "root": str(c.root()),
        "remember_scope": remembered.get("scope") or "campaign",
        **excel,
    }


_CELL_REF_RE = re.compile(r"^(?:(?P<sheet>.+)!)?(?P<cell>[A-Za-z]{1,3}\d{1,5})$")


def _split_cell_ref(spec: str) -> tuple[str, str]:
    raw = str(spec or "").strip()
    m = _CELL_REF_RE.match(raw)
    if not m:
        return "", ""
    return (m.group("sheet") or "").strip("'"), m.group("cell").upper()


def _row_label_is_tags(ws, cell: str) -> bool:
    from openpyxl.utils.cell import coordinate_from_string

    _col, row = coordinate_from_string(cell)
    lab = str(ws.cell(row, 1).value or "").strip().lower().rstrip(":")
    return lab in ("tags", "tag", "labels", "campaign tags")


def _find_tags_cell(wb, spec: str) -> tuple[str, str, bool]:
    sheet, cell = _split_cell_ref(spec)
    if cell:
        if sheet and sheet in wb.sheetnames and _row_label_is_tags(wb[sheet], cell):
            return sheet, cell, False
        if (not sheet) and "Summary" in wb.sheetnames and _row_label_is_tags(wb["Summary"], cell):
            return "Summary", cell, False
    names: list[str] = []
    if "Summary" in wb.sheetnames:
        names.append("Summary")
    names.extend(n for n in wb.sheetnames if n not in names)
    for name in names:
        ws = wb[name]
        for r in range(1, 61):
            lab = str(ws.cell(r, 1).value or "").strip().lower().rstrip(":")
            if lab in ("tags", "tag", "labels", "campaign tags"):
                return name, f"B{r}", False
    sh = "Summary" if "Summary" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sh]
    r = int(ws.max_row or 1) + 2
    ws.cell(r, 1).value = "Tags"
    return sh, f"B{r}", True


def _remember_tags_cell(ctx, sheet_cell: str) -> None:
    path = ctx.sheet_map_path()
    if not path.is_file():
        return
    sm = ctx.load_sheet_map() or {}
    ident = sm.get("identity")
    if isinstance(ident, dict) and str(ident.get("tags") or "").strip():
        return
    text = path.read_text(encoding="utf-8")
    if re.search(r"(?m)^identity:\s*$", text):
        path.write_text(
            re.sub(r"(?m)^(identity:\s*)$", rf"\1\n  tags: {sheet_cell}", text, count=1),
            encoding="utf-8",
        )
        return
    path.write_text(text.rstrip() + f"\nidentity:\n  tags: {sheet_cell}\n", encoding="utf-8")


def sync_tags_to_workbook(ctx, tags: list[str]) -> dict[str, Any]:
    """Stamp joined tag tokens into the lab xlsx Tags cell. Never fails save_tags."""
    path = ctx.lab_report_path()
    if not path.is_file():
        return {"excel": None, "excel_status": "no_workbook"}
    sm = ctx.load_sheet_map() or {}
    ident = sm.get("identity") if isinstance(sm.get("identity"), dict) else {}
    spec = str((ident or {}).get("tags") or "").strip()
    wb = None
    try:
        from openpyxl import load_workbook

        wb = load_workbook(path)
        if not wb.sheetnames:
            return {"excel": str(path), "excel_status": "no_sheet"}
        sheet, cell, created = _find_tags_cell(wb, spec)
        wb[sheet][cell] = ", ".join(tags)
        wb.save(path)
        sheet_cell = f"{sheet}!{cell}"
        try:
            _remember_tags_cell(ctx, sheet_cell)
        except Exception:
            pass
        return {
            "excel": str(path),
            "excel_cell": sheet_cell,
            "excel_status": "ok",
            "excel_created_label": created,
        }
    except PermissionError as exc:
        return {"excel": str(path), "excel_status": "locked", "excel_error": str(exc)}
    except OSError as exc:
        msg = str(exc).lower()
        locked = getattr(exc, "errno", None) in (13, 5) or "denied" in msg or "being used" in msg
        return {
            "excel": str(path),
            "excel_status": "locked" if locked else "error",
            "excel_error": str(exc),
        }
    except Exception as exc:
        return {"excel": str(path), "excel_status": "error", "excel_error": str(exc)}
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass


def _load_boards_yaml() -> dict[str, Any]:
    if not BOARDS_PATH.is_file():
        return {}
    data = yaml.safe_load(BOARDS_PATH.read_text(encoding="utf-8")) or {}
    return data if isinstance(data, dict) else {}


def load_stored_vocab() -> dict[str, Any]:
    if not LABEL_VOCAB_PATH.is_file():
        return {"families": {}, "all": []}
    data = yaml.safe_load(LABEL_VOCAB_PATH.read_text(encoding="utf-8")) or {}
    if not isinstance(data, dict):
        return {"families": {}, "all": []}
    fams = data.get("families") if isinstance(data.get("families"), dict) else {}
    all_labs = _normalize_labels(data.get("all") or [])
    clean_fams: dict[str, list[dict[str, str]]] = {}
    for key, rows in fams.items():
        fam = _family_key(str(key))
        if fam:
            clean_fams[fam] = _normalize_labels(rows)
    return {"families": clean_fams, "all": all_labs}


def remember_labels(
    labels: list[dict[str, str]] | None,
    *,
    scope: str = "campaign",
    family: str = "",
    ctx=None,
) -> dict[str, Any]:
    """Persist class/all vocab. Campaign tokens already live in tags.yaml."""
    from ate.core.database import family_for_component, get_context

    want = str(scope or "campaign").strip().lower()
    if want in ("class", "this class", "this_class"):
        want = "family"
    if want in ("all products", "all_products", "global"):
        want = "all"
    if want not in ("family", "all"):
        return {"scope": "campaign"}
    parsed = _normalize_labels(labels or [])
    if not parsed:
        return {"scope": want, "added": 0}
    c = ctx or get_context()
    fam = _family_key(family or family_for_component(getattr(c, "component", "") or "") or "")
    stored = load_stored_vocab()
    if want == "all":
        merged = _normalize_labels(list(stored.get("all") or []) + parsed)
        stored["all"] = merged
    else:
        if not fam:
            return {"scope": "campaign", "reason": "no family"}
        fams = stored.setdefault("families", {})
        fams[fam] = _normalize_labels(list(fams.get(fam) or []) + parsed)
    payload = {
        "families": stored.get("families") or {},
        "all": stored.get("all") or [],
    }
    LABEL_VOCAB_PATH.write_text(
        _LABEL_VOCAB_PREAMBLE + yaml.safe_dump(payload, sort_keys=False, allow_unicode=True),
        encoding="utf-8",
    )
    return {"scope": want, "family": fam, "added": len(parsed)}


def _family_key(family: str) -> str:
    fam = (family or "").strip().lower()
    if fam in ("analog_switch", "lim"):
        return "switch"
    return fam


def list_kinds(*, family: str = "", ctx=None) -> list[dict[str, str]]:
    data = _load_boards_yaml()
    kinds_block = data.get("kinds") if isinstance(data.get("kinds"), dict) else {}
    out: dict[str, str] = {}
    for kid, meta in kinds_block.items():
        k = _norm_kind(str(kid))
        label = str((meta or {}).get("label") or k.replace("_", " ")) if isinstance(meta, dict) else k
        out[k] = label
    if "board" not in out:
        out["board"] = "Board"
    if "board_type" not in out:
        out["board_type"] = "Board type"
    stored = load_stored_vocab()
    extra_labs: list[dict[str, str]] = list(stored.get("all") or [])
    fam = _family_key(family)
    if fam:
        extra_labs.extend((stored.get("families") or {}).get(fam) or [])
    if ctx is None:
        try:
            from ate.core.database import get_context

            ctx = get_context()
        except Exception:
            ctx = None
    if ctx is not None:
        extra_labs.extend(list(load_tags(ctx).get("labels") or []))
    for lab in extra_labs:
        k = str(lab.get("kind") or "")
        if k and k not in out:
            out[k] = k.replace("_", " ")
    return [{"id": k, "label": lab} for k, lab in out.items()]


def list_boards(
    *,
    family: str = "",
    package: str = "",
    ctx=None,
) -> list[str]:
    from ate.core.database import family_for_component, get_context

    c = ctx or get_context()
    fam = _family_key(family or family_for_component(c.component) or "")
    pkg = (package or c.package or "").strip()
    if not fam:
        return []
    families = _load_boards_yaml().get("families") or {}
    fam_block = families.get(fam) or {}
    if not isinstance(fam_block, dict):
        return []
    rows = fam_block.get(pkg) if pkg and pkg not in _SKIP_FAM_KEYS else None
    if not isinstance(rows, list):
        rows = fam_block.get("default") or []
    if not isinstance(rows, list):
        return []
    return [str(x) for x in rows if str(x).strip()]


def _kind_values_from_yaml(fam: str, kind: str, package: str) -> list[str]:
    families = _load_boards_yaml().get("families") or {}
    fam_block = families.get(fam) or {}
    if not isinstance(fam_block, dict):
        return []
    if kind == "board":
        return list_boards(family=fam, package=package)
    key = "board_types" if kind == "board_type" else kind
    rows = fam_block.get(key)
    if not isinstance(rows, list):
        rows = fam_block.get(f"{kind}s") or []
    if not isinstance(rows, list):
        return []
    return [str(x) for x in rows if str(x).strip()]


def _scan_labels(component: str = "") -> list[dict[str, str]]:
    """Reuse labels from campaigns. No Tags folder axis."""
    from ate.core import paths as pathmod

    root = pathmod.TEST_DB_ROOT
    if component:
        root = root / component
    if not root.is_dir():
        return []
    found: list[dict[str, str]] = []
    # ponytail: rglob tags.yaml; add an index file if this walk exceeds ~2s
    for yml in root.rglob("tags.yaml"):
        if yml.parent.name != "_manifest":
            continue
        try:
            data = yaml.safe_load(yml.read_text(encoding="utf-8")) or {}
        except OSError:
            continue
        if isinstance(data, dict):
            found.extend(_labels_from_stored(data))
    return found


def _scan_family_labels(component: str) -> list[dict[str, str]]:
    return _scan_labels(component)


def _merge_kind_values(values: dict[str, list[str]], labels: list[dict[str, str]]) -> None:
    for lab in labels:
        k = str(lab.get("kind") or "")
        v = str(lab.get("value") or "")
        if not k or not v:
            continue
        values.setdefault(k, [])
        if v not in values[k]:
            values[k].append(v)


def list_label_vocab(
    *,
    family: str = "",
    package: str = "",
    ctx=None,
    scan: str = "all",
) -> dict[str, Any]:
    from ate.core.database import family_for_component, get_context

    c = ctx or get_context()
    fam = _family_key(family or family_for_component(c.component) or "")
    pkg = (package or c.package or "").strip()
    values: dict[str, list[str]] = {}
    for kind in ("board", "board_type"):
        values[kind] = _kind_values_from_yaml(fam, kind, pkg)
    stored = load_stored_vocab()
    campaign = load_tags(c) if c else {"labels": []}
    _merge_kind_values(values, list(campaign.get("labels") or []))
    _merge_kind_values(values, list((stored.get("families") or {}).get(fam) or []))
    _merge_kind_values(values, list(stored.get("all") or []))
    want = str(scan or "all").strip().lower()
    if want in ("family", "all"):
        _merge_kind_values(values, _scan_labels(getattr(c, "component", "") or ""))
    if want == "all":
        _merge_kind_values(values, _scan_labels(""))
    return {
        "family": fam,
        "values": values,
        "labels": list(campaign.get("labels") or []),
        "scopes": ["campaign", "family", "all"],
    }


def list_label_catalog(*, family: str = "", package: str = "", ctx=None) -> dict[str, Any]:
    vocab = list_label_vocab(family=family, package=package, ctx=ctx)
    return {
        "kinds": list_kinds(family=family, ctx=ctx),
        "values": vocab.get("values") or {},
        "labels": vocab.get("labels") or [],
        "family": vocab.get("family") or "",
    }


def import_tags(from_root: str | Path, *, ctx=None, merge: bool = True) -> dict[str, Any]:
    """Copy tags from another campaign root's tags.yaml into the active campaign."""
    from ate.core.database import get_context

    c = ctx or get_context()
    src = Path(from_root)
    if src.is_dir():
        yml = src / "_manifest" / "tags.yaml"
        if not yml.is_file():
            yml = src / "tags.yaml"
    else:
        yml = src
    if not yml.is_file():
        raise FileNotFoundError(f"No tags.yaml at {from_root}")
    data = yaml.safe_load(yml.read_text(encoding="utf-8")) or {}
    if not isinstance(data, dict):
        raise ValueError("tags.yaml must be a mapping")
    incoming = _labels_from_stored(data)
    if merge:
        cur = load_tags(c)
        incoming = list(cur.get("labels") or []) + incoming
    return save_tags(labels=incoming, ctx=c)


def filter_campaigns_by_tag(
    tag: str,
    *,
    component: str = "",
) -> list[dict[str, Any]]:
    """Scan #Test_Database for campaigns whose TAGS.txt / tags.yaml contain tag."""
    from ate.core import paths as pathmod

    want = str(tag or "").strip().lower()
    if not want:
        return []
    root = pathmod.TEST_DB_ROOT
    if not root.is_dir():
        return []
    hits: list[dict[str, Any]] = []
    for tags_file in root.rglob("TAGS.txt"):
        try:
            text = tags_file.read_text(encoding="utf-8").lower()
        except OSError:
            continue
        if want not in text:
            continue
        campaign = tags_file.parent
        # Expect .../Component/Part/Package/Operator/Version_N/TAGS.txt
        parts = campaign.parts
        try:
            idx = next(
                i
                for i, x in enumerate(parts)
                if x in ("#Test_Database", "Test_Database")
            )
            comp = parts[idx + 1]
            part = parts[idx + 2]
            package = parts[idx + 3]
            operator = parts[idx + 4]
            version = parts[idx + 5]
        except (StopIteration, IndexError):
            continue
        if component and comp.lower() != component.strip().lower():
            continue
        hits.append(
            {
                "component": comp,
                "part": part,
                "package": package,
                "operator": operator,
                "version": version,
                "root": str(campaign),
                "tags_txt": str(tags_file),
            }
        )
    return hits


def list_tags_rpc(*, ctx=None) -> dict[str, Any]:
    from ate.core.database import get_context

    c = ctx or get_context()
    data = load_tags(c)
    cat = list_label_catalog(ctx=c)
    return {
        **data,
        "boards_vocab": list_boards(ctx=c),
        "kinds": cat.get("kinds") or [],
        "label_values": cat.get("values") or {},
        "tags_yaml": str(tags_yaml_path(c)),
        "tags_txt": str(tags_txt_path(c)),
        "root": str(c.root()),
        "family": cat.get("family") or "",
        "scopes": ["campaign", "family", "all"],
    }
