"""Photo-grid SoT check (A07-T01).

Run: python -m ate.reporting.check_photo_layout
"""
from __future__ import annotations

import sys
from pathlib import Path

from ate.reporting import lab_report, photo_layout
from ate.core import campaign_outline as campaign_outline_mod


_BANNED = (
    "_SSSR_PHOTO_ANCHORS",
    "_LSSR_PHOTO_ANCHORS",
    "_SETTLING_UNIT_COLS",
    "_SETTLING_BODY_ROW",
)


def _sample_yaml() -> str:
    return (
        "tests:\n"
        "  Foo:\n"
        "    paste:\n"
        "      photos:\n"
        "        u1_chA: A1  # keep comment\n"
        "        u1_chB: E1\n"
        "  Bar:\n"
        "    paste:\n"
        "      photos:\n"
        "        u1_chA: B1\n"
    )


def check_photo_layout() -> list[str]:
    errors: list[str] = []
    src = Path(lab_report.__file__).read_text(encoding="utf-8")
    for name in _BANNED:
        if name in src:
            errors.append(f"lab_report.py still hardcodes {name}")

    parsed = photo_layout.parse_photo_key("u2_chB")
    if not parsed or parsed["key"] != "u2_chB" or parsed["unit"] != 2:
        errors.append("parse_photo_key u2_chB failed")
    ort = photo_layout.parse_photo_key("pos_u1_chA")
    if not ort or ort["key"] != "pos_u1_chA" or ort["polarity"] != "POS":
        errors.append("parse_photo_key pos_u1_chA failed")
    trial = photo_layout.parse_photo_key("t2_u1_chA")
    if not trial or trial["key"] != "t2_u1_chA" or trial.get("trial") != 2:
        errors.append("parse_photo_key t2_u1_chA failed")

    outline_src = Path(campaign_outline_mod.__file__).read_text(encoding="utf-8")
    if "_PHOTO_TTSOP" in outline_src:
        errors.append("campaign_outline.py still hardcodes _PHOTO_TTSOP")

    from openpyxl import Workbook
    from ate.reporting.golden_layout import (
        apply_photo_bands,
        discover_photo_anchors,
        ensure_photo_boxes,
        photo_header_start,
    )

    wb = Workbook()
    short = wb.active
    short.title = "Short"
    short["A5"] = "Test Conclusion"
    start_s = photo_header_start(short)
    a_short = apply_photo_bands(short, start_row=start_s, sample_size=2, titles=["Slew"])
    if set(a_short) != {"u1_chA", "u1_chB", "u2_chA", "u2_chB"}:
        errors.append(f"sample_size=2 should be 4 DUT/ch boxes, got {a_short}")
    if discover_photo_anchors(short) != a_short:
        errors.append("discover_photo_anchors must match what was placed")
    heights = {
        mr.max_row - mr.min_row + 1
        for mr in short.merged_cells.ranges
        if mr.max_col - mr.min_col + 1 == 4 and mr.max_row - mr.min_row + 1 >= 6
    }
    if heights != {10}:
        errors.append(f"photo body must be 10 rows, got {heights}")

    long = wb.create_sheet("Long")
    long["A20"] = "Test Conclusion"
    start_l = photo_header_start(long)
    if start_l <= start_s:
        errors.append(f"longer intro must push photo start down ({start_l} vs {start_s})")
    a_long = apply_photo_bands(
        long, start_row=start_l, sample_size=4, titles=["Positive X", "Negative X"]
    )
    if len(a_long) != 16:
        errors.append(f"4 DUT x POS/NEG should be 16 boxes, got {len(a_long)}")
    if "pos_u1_chA" not in a_long or "neg_u4_chB" not in a_long:
        errors.append(f"polarity keys missing: {sorted(a_long)}")

    trials = wb.create_sheet("Trials")
    trials["A4"] = "Test Conclusion"
    tanch = apply_photo_bands(
        trials,
        start_row=photo_header_start(trials),
        sample_size=1,
        titles=["1VPP", "2VPP", "NEG"],
    )
    if set(tanch) != {"u1_chA", "u1_chB", "t2_u1_chA", "t2_u1_chB", "t3_u1_chA", "t3_u1_chB"}:
        errors.append(f"trial bands should stack t2/t3, got {tanch}")

    empty = wb.create_sheet("Empty")
    empty["A3"] = "Test Conclusion"
    if ensure_photo_boxes(empty, sample_size=4, write=False):
        errors.append("ensure_photo_boxes(write=False) must not invent boxes")
    kept = ensure_photo_boxes(short, write=False)
    if kept != a_short:
        errors.append("ensure_photo_boxes must reuse existing merges")

    wide = wb.create_sheet("Wide")
    wide["A5"] = "Test Conclusion"
    wide.merge_cells("A8:H17")
    wide.merge_cells("I8:P17")
    wide_found = discover_photo_anchors(wide)
    if wide_found.get("u1_chA") != "A8" or wide_found.get("u1_chB") != "I8":
        errors.append(f"8-col CHA|CHB pair must be u1_chA/u1_chB, got {wide_found}")
    if "u2_chA" in wide_found:
        errors.append("two 8-col boxes must not map as DUT2")

    fall = wb.create_sheet("Fall")
    fall["A5"] = "Test Conclusion"
    fall["A6"] = "Positive slew"
    fall["A20"] = "Negative slew"
    fall.merge_cells("A8:H17")
    fall.merge_cells("I8:P17")
    fall.merge_cells("A22:H31")
    fall.merge_cells("I22:P31")
    fall_found = discover_photo_anchors(fall)
    if "pos_u1_chA" not in fall_found or "neg_u1_chB" not in fall_found:
        errors.append(f"pos/neg 8-col bands missing: {sorted(fall_found)}")

    patched = photo_layout.patch_photos_yaml(
        _sample_yaml(), "Foo", {"u1_chA": "z9"}
    )
    if "u1_chA: Z9" not in patched:
        errors.append("YAML patch did not update Foo u1_chA")
    if "# keep comment" not in patched:
        errors.append("YAML patch dropped comment")
    if "u1_chA: B1" not in patched:
        errors.append("YAML patch mutated Bar")

    try:
        photo_layout.patch_photos_yaml(_sample_yaml(), "Foo", {"u1_chA": "not_a_cell"})
        errors.append("invalid cell was accepted")
    except ValueError:
        pass

    try:
        photo_layout.resolve_shot_file("../secret.png")
        errors.append("path traversal accepted")
    except (PermissionError, FileNotFoundError):
        pass

    try:
        tests = photo_layout.list_photo_tests()
    except Exception as exc:
        errors.append(f"list_photo_tests: {exc}")
        return errors
    if tests:
        key = tests[0]["test_key"]
        try:
            grid = photo_layout.photos_map(key)
        except Exception as exc:
            errors.append(f"photos_map({key}): {exc}")
            return errors
        if not grid:
            errors.append(f"{key} paste.photos empty after parse")
        else:
            first = next(iter(grid))
            parsed_first = photo_layout.parse_photo_key(first)
            if parsed_first:
                got = photo_layout.photo_anchor(
                    key,
                    parsed_first["unit"],
                    parsed_first["channel"],
                    parsed_first["polarity"],
                )
                if got != grid[first]:
                    errors.append(f"photo_anchor drift {key} {first}: {got} vs {grid[first]}")
    return errors


def main() -> int:
    errors = check_photo_layout()
    if errors:
        print("FAIL photo-layout:")
        for line in errors:
            print(f"  - {line}")
        return 1
    print("OK photo-layout: merged-box discover, 8-col CHA/CHB, no _PHOTO_TTSOP, YAML patch keeps comments")
    return 0


if __name__ == "__main__":
    sys.exit(main())
