"""Golden ORT-style layout: wrap-text merge sizing + scale to all lab sheets.

Culture (do not drag-resize):
  1) Count wrap lines / chars → compute how many ROWS and COLS to merge
  2) Merge & center (or top-left wrap) that block
  3) Zones: left intro | right-top precautions+results | bottom photo boxes
  4) sample_size = 4 → DUT #1..#4 photo slots

Non-ORT sheets: insert_rows to expand intro to wrap budget (pushes tables down),
append ORT-geometry photo grid below content.
"""
from __future__ import annotations

import math
import re
from dataclasses import dataclass
from typing import Any

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from ate.reporting.sheet_layout import (
    HEADER_BLUE,
    HEADER_BLUE_TEMPLATE,
    ORT_LAYOUT,
    YELLOW,
    _apply_border_range,
    _ensure_merge,
    _fill,
    _font_white_bold,
    apply_standard_column_widths,
    ensure_sample_size,
)

TEST_SHEETS = [
    "Slew Rate",
    "GBW",
    "SettlingTime",
    "SSSR",
    "LSSR",
    "ORT",
    "PSRR",
    "CMRR",
    "AOL",
    "EMIRR",
    "NoPhaseReversal",
    "VOS",
    "PowerOnTime",
    "VOL",
    "Noise",
]

PHOTO_TITLE_BY_SHEET = {
    "Slew Rate": "Slew Rate",
    "GBW": "GBW",
    "SettlingTime": "Settling Time",
    "SSSR": "Small Signal Step",
    "LSSR": "Large Signal Step",
    "ORT": "Overload Recovery Time",
    "PSRR": "PSRR",
    "CMRR": "CMRR",
    "AOL": "AOL",
    "EMIRR": "EMIRR",
    "NoPhaseReversal": "Phase Reversal",
    "VOS": "VOS",
    "PowerOnTime": "Power On Time",
    "VOL": "VOHL / VOL",
    "Noise": "1-10Hz Noise",
}

ERROR_TOKENS = {"#VALUE!", "#REF!", "#N/A", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!"}


@dataclass
class WrapBudget:
    text: str
    merge_cols: int
    rows: int
    lines: int
    chars_per_line: int


def normalize_breaks(text: Any) -> str:
    if text is None:
        return ""
    s = str(text)
    if s.strip().upper() in ERROR_TOKENS:
        return ""
    s = s.replace("\r\n", "\n").replace("\r", "\n")
    s = re.sub(r"<br\s*/?>", "\n", s, flags=re.I)
    s = s.replace("\\n", "\n")
    return s


def chars_per_line(merge_cols: int, col_width: float = 16.0) -> int:
    return max(8, int(merge_cols * col_width * 0.95))


def count_wrap_lines(text: str, cpl: int) -> int:
    if not text.strip():
        return 1
    total = 0
    for para in text.split("\n"):
        if not para:
            total += 1
            continue
        total += max(1, math.ceil(len(para) / cpl))
    return max(1, total)


def rows_for_wrap(
    text: Any,
    *,
    merge_cols: int = 11,
    col_width: float = 16.0,
    min_rows: int = 2,
    max_rows: int = 20,
    pad_rows: int = 1,
) -> WrapBudget:
    clean = normalize_breaks(text)
    cpl = chars_per_line(merge_cols, col_width)
    lines = count_wrap_lines(clean, cpl)
    rows = min(max_rows, max(min_rows, lines + pad_rows))
    return WrapBudget(text=clean, merge_cols=merge_cols, rows=rows, lines=lines, chars_per_line=cpl)


def clear_error_placeholders(ws, max_row: int = 200, max_col: int = 40) -> int:
    n = 0
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for cell in row:
            try:
                if cell.value is not None and str(cell.value).strip().upper() in ERROR_TOKENS:
                    cell.value = None
                    n += 1
            except AttributeError:
                pass
    return n


def _style_label(cell, fill_rgb: str = HEADER_BLUE_TEMPLATE) -> None:
    cell.fill = _fill(fill_rgb)
    cell.font = Font(bold=True, color="FFFFFF", size=11)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_wrap_value(cell) -> None:
    # Golden rule: everything middle + center
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def hard_expand_intro_with_insert(
    ws,
    label_row: int,
    *,
    value_end_col: int = 12,
    min_rows: int = 2,
    max_rows: int = 16,
) -> dict:
    """Insert rows so wrap-text merge gets needed height; pushes content below down."""
    lab = ws.cell(label_row, 1).value
    val = ws.cell(label_row, 2).value
    if not lab:
        return {}

    budget = rows_for_wrap(
        val, merge_cols=value_end_col - 1, min_rows=min_rows, max_rows=max_rows
    )
    key = str(lab).strip().lower()
    if key == "test circuitry":
        budget = WrapBudget(
            text=budget.text,
            merge_cols=budget.merge_cols,
            rows=13,  # fixed golden budget
            lines=budget.lines,
            chars_per_line=budget.chars_per_line,
        )
    elif key == "datasheet":
        budget = WrapBudget(
            text=budget.text,
            merge_cols=budget.merge_cols,
            rows=7,  # fixed golden budget
            lines=budget.lines,
            chars_per_line=budget.chars_per_line,
        )
    elif key == "test conclusion":
        budget = WrapBudget(
            text=budget.text,
            merge_cols=budget.merge_cols,
            rows=7,  # fixed golden budget
            lines=budget.lines,
            chars_per_line=budget.chars_per_line,
        )

    current_span = 1
    for mr in ws.merged_cells.ranges:
        if mr.min_row == label_row and mr.min_col == 2:
            current_span = mr.max_row - mr.min_row + 1
            break

    need = budget.rows
    if need > current_span:
        # Unmerge anything that would be split by insert_rows (openpyxl can drop labels)
        from openpyxl.utils import range_boundaries

        insert_at = label_row + current_span
        for mr in list(ws.merged_cells.ranges):
            rmin, rmax = mr.min_row, mr.max_row
            if rmax >= insert_at and rmin < insert_at + (need - current_span):
                try:
                    ws.unmerge_cells(str(mr))
                except Exception:
                    pass
        ws.insert_rows(insert_at, amount=need - current_span)

    end_row = label_row + need - 1

    # Clear body of the block before merge (keep top-left values)
    for r in range(label_row, end_row + 1):
        for c in range(1, value_end_col + 1):
            if r == label_row and c in (1, 2):
                continue
            try:
                ws.cell(r, c).value = None
            except AttributeError:
                pass

    _ensure_merge(ws, f"A{label_row}:A{end_row}")
    _ensure_merge(ws, f"B{label_row}:{get_column_letter(value_end_col)}{end_row}")

    lab_cell = ws.cell(label_row, 1)
    lab_cell.value = str(lab)
    _style_label(lab_cell)
    val_cell = ws.cell(label_row, 2)
    if budget.text:
        val_cell.value = budget.text
    elif val is not None and str(val).strip().upper() in ERROR_TOKENS:
        val_cell.value = None
    _style_wrap_value(val_cell)
    # Borders applied later for full A1:L{conclusion} box — avoid per-block spray
    for r in range(label_row, end_row + 1):
        ws.row_dimensions[r].height = 20.4

    return {
        "label": str(lab),
        "start": label_row,
        "end": end_row,
        "wanted_rows": need,
        "used_rows": need,
        "lines": budget.lines,
        "chars_per_line": budget.chars_per_line,
        "inserted": max(0, need - current_span),
    }


def apply_photo_grid_single(
    ws,
    *,
    title: str,
    start_row: int,
    sample_size: int = 4,
    body_rows: int = 10,
) -> dict[str, str]:
    """Bottom photo grid: sample_size × (ChA, ChB). Golden ORT geometry."""
    yellow = _fill(YELLOW)
    blue = _fill(HEADER_BLUE)
    wrap_center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    box_starts = (1, 5, 9, 13, 17, 21, 25, 29)[: sample_size * 2]
    unit_starts = (1, 9, 17, 25)[:sample_size]
    header_row = start_row
    channel_row = start_row + 1
    body0 = start_row + 2
    body1 = body0 + body_rows - 1

    for i, start in enumerate(unit_starts, start=1):
        end = start + 7
        _ensure_merge(
            ws, f"{get_column_letter(start)}{header_row}:{get_column_letter(end)}{header_row}"
        )
        cell = ws.cell(header_row, start)
        cell.value = f"{title} #{i}"
        cell.fill = blue
        cell.font = _font_white_bold()
        cell.alignment = wrap_center

    anchors: dict[str, str] = {}
    channels = ("Channel A", "Channel B")
    for box_i, start_col in enumerate(box_starts):
        end_col = start_col + 3
        unit = box_i // 2 + 1
        ch_key = "chA" if box_i % 2 == 0 else "chB"

        _ensure_merge(
            ws,
            f"{get_column_letter(start_col)}{channel_row}:{get_column_letter(end_col)}{channel_row}",
        )
        c = ws.cell(channel_row, start_col)
        c.value = channels[box_i % 2]
        c.fill = yellow
        c.font = Font(bold=True)
        c.alignment = wrap_center
        _apply_border_range(ws, channel_row, channel_row, start_col, end_col)

        for r in range(body0, body1 + 1):
            for col in range(start_col, end_col + 1):
                cell = ws.cell(r, col)
                try:
                    if cell.value is not None and str(cell.value).strip().upper() in ERROR_TOKENS:
                        cell.value = None
                    elif (r != body0 or col != start_col) and cell.value in (None, ""):
                        cell.value = None
                except AttributeError:
                    pass
                cell.alignment = wrap_center
        try:
            ws.cell(body0, start_col).value = None
        except AttributeError:
            pass

        _ensure_merge(
            ws,
            f"{get_column_letter(start_col)}{body0}:{get_column_letter(end_col)}{body1}",
        )
        _apply_border_range(ws, body0, body1, start_col, end_col)
        anchors[f"u{unit}_{ch_key}"] = f"{get_column_letter(start_col)}{body0}"

    for r in range(header_row, body1 + 1):
        ws.row_dimensions[r].height = 20.4
    return anchors


PHOTO_START_COLS = (1, 5, 9, 13, 17, 21, 25, 29)
PHOTO_COLS_PER_BOX = 4
PHOTO_BODY_ROWS = 10
PHOTO_BODY_MIN = 6
PHOTO_BODY_MAX = 16
PHOTO_BAND_GAP = 1
# 4-col = DUT x CHA/CHB grid. 8-col = OpAmp golden CHA | CHB pair (Slew Rate).
PHOTO_START_COLS_BY_WIDTH = {
    4: PHOTO_START_COLS,
    8: (1, 9, 17, 25),
}


def _photo_body_merges(ws) -> list:
    out = []
    for mr in ws.merged_cells.ranges:
        width = mr.max_col - mr.min_col + 1
        height = mr.max_row - mr.min_row + 1
        starts = PHOTO_START_COLS_BY_WIDTH.get(width)
        if (
            starts
            and PHOTO_BODY_MIN <= height <= PHOTO_BODY_MAX
            and mr.min_col in starts
        ):
            out.append(mr)
    out.sort(key=lambda m: (m.min_row, m.min_col))
    return out


def intro_end_row(ws, max_scan: int = 80) -> int:
    """Last row of the Test Conclusion merge (or the label row)."""
    for row in range(1, max_scan + 1):
        lab = ws.cell(row, 1).value
        if not lab or str(lab).strip().lower() != "test conclusion":
            continue
        end = row
        for mr in ws.merged_cells.ranges:
            if mr.min_row == row and mr.min_col <= 2:
                end = max(end, mr.max_row)
        return end
    return 1


def photo_header_start(ws, *, gap: int = 2) -> int:
    """Header row for the first photo band: existing boxes, else after conclusion."""
    bodies = _photo_body_merges(ws)
    if bodies:
        return max(1, min(m.min_row for m in bodies) - 2)
    return intro_end_row(ws) + max(1, int(gap))


def _col_unit_ch(col: int, width: int = 4) -> tuple[int, str]:
    starts = PHOTO_START_COLS_BY_WIDTH.get(int(width) or 4) or PHOTO_START_COLS
    idx = list(starts).index(int(col)) if int(col) in starts else PHOTO_START_COLS.index(int(col))
    unit = idx // 2 + 1
    ch = "chA" if idx % 2 == 0 else "chB"
    return unit, ch


def _band_key_prefix(titles: list[str], index: int) -> str:
    lows = [str(t or "").lower() for t in titles]
    this = lows[index] if 0 <= index < len(lows) else ""
    if "positive" in this:
        return "pos_"
    if "negative" in this:
        return "neg_"
    pos = [i for i, t in enumerate(lows) if "positive" in t]
    neg = [i for i, t in enumerate(lows) if "negative" in t]
    if len(titles) == 2 and len(pos) == 1 and len(neg) == 1:
        return "pos_" if index in pos else "neg_"
    if index == 0:
        return ""
    return f"t{index + 1}_"


def discover_photo_anchors(ws) -> dict[str, str]:
    """Top-left of each merged 4-col or 8-col photo body.

    Two boxes on a row = Channel A | Channel B (OpAmp golden Slew), not DUT2.
    Extra bands: pos_/neg_ from titles, else tN_.
    """
    bodies = _photo_body_merges(ws)
    if not bodies:
        return {}
    rows = sorted({m.min_row for m in bodies})
    titles: list[str] = []
    for row in rows:
        header = ws.cell(row - 2, 1).value or ws.cell(row - 1, 1).value or ""
        titles.append(str(header))
    anchors: dict[str, str] = {}
    for band_i, row in enumerate(rows):
        prefix = _band_key_prefix(titles, band_i)
        band = [mr for mr in bodies if mr.min_row == row]
        band.sort(key=lambda m: m.min_col)
        if len(band) == 2:
            pairs = [(1, "chA"), (1, "chB")]
        else:
            pairs = [
                _col_unit_ch(mr.min_col, width=mr.max_col - mr.min_col + 1) for mr in band
            ]
        for mr, (unit, ch) in zip(band, pairs):
            anchors[f"{prefix}u{unit}_{ch}"] = f"{get_column_letter(mr.min_col)}{mr.min_row}"
    return anchors


def apply_photo_bands(
    ws,
    *,
    start_row: int,
    sample_size: int = 4,
    titles: list[str] | None = None,
    body_rows: int = PHOTO_BODY_ROWS,
) -> dict[str, str]:
    """Stack same-size merged bands. sample_size DUT columns; len(titles) = trial count."""
    titles = list(titles or ["DUT"])
    sample = max(1, int(sample_size or 1))
    anchors: dict[str, str] = {}
    row = int(start_row)
    for i, title in enumerate(titles):
        band = apply_photo_grid_single(
            ws, title=str(title), start_row=row, sample_size=sample, body_rows=body_rows
        )
        prefix = _band_key_prefix(titles, i)
        for key, cell in band.items():
            anchors[f"{prefix}{key}"] = cell
        row += 2 + int(body_rows) + PHOTO_BAND_GAP
    return anchors


def ensure_photo_boxes(
    ws,
    *,
    sample_size: int = 4,
    titles: list[str] | None = None,
    write: bool = False,
) -> dict[str, str]:
    """Reuse existing merged boxes. Place only when write=True and none exist."""
    found = discover_photo_anchors(ws)
    if found:
        return found
    if not write:
        return {}
    return apply_photo_bands(
        ws,
        start_row=photo_header_start(ws),
        sample_size=sample_size,
        titles=titles,
    )


def find_content_bottom(ws, max_scan: int = 180) -> int:
    bottom = 1
    for row in ws.iter_rows(min_row=1, max_row=max_scan, max_col=16):
        for cell in row:
            if cell.value not in (None, ""):
                bottom = max(bottom, cell.row)
    return bottom


def ensure_right_precautions(ws) -> None:
    from ate.reporting.rules_layout import ensure_precautions_n1q4

    ensure_precautions_n1q4(ws)


def ensure_sample_columns_headers(ws, sample_size: int = 4) -> int:
    touched = 0
    for row in range(1, 120):
        for col in range(1, 40):
            val = ws.cell(row, col).value
            if val is None:
                continue
            # Only "#1" — never bare "1" (that corrupted Precautions N2)
            if str(val).strip() == "#1":
                for i in range(sample_size):
                    cell = ws.cell(row, col + i)
                    try:
                        cell.value = f"#{i + 1}"
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.font = Font(bold=True)
                        touched += 1
                    except AttributeError:
                        pass
                break
    return touched


def apply_golden_non_ort(ws, sheet_name: str, sample_size: int = 4) -> dict:
    cleared = clear_error_placeholders(ws)
    ensure_sample_size(ws, sample_size)
    apply_standard_column_widths(ws, ORT_LAYOUT)
    ensure_right_precautions(ws)

    intro_reports = []
    targets = (
        ("test conditions", 4),
        ("test circuitry", 13),
        ("datasheet", 7),
        ("test conclusion", 7),
    )
    for key, mins in targets:
        label_row = None
        for row in range(1, 80):
            lab = ws.cell(row, 1).value
            if lab and str(lab).strip().lower() == key:
                label_row = row
                break
        if label_row is None:
            continue
        intro_reports.append(
            hard_expand_intro_with_insert(ws, label_row, min_rows=mins, max_rows=16)
        )

    sample_hdrs = ensure_sample_columns_headers(ws, sample_size)

    title = PHOTO_TITLE_BY_SHEET.get(sheet_name, sheet_name)
    existing = discover_photo_anchors(ws)
    if existing:
        anchors = existing
        photo_start = photo_header_start(ws)
    else:
        photo_start = photo_header_start(ws)
        anchors = apply_photo_bands(
            ws,
            start_row=photo_start,
            sample_size=sample_size,
            titles=[title],
            body_rows=PHOTO_BODY_ROWS,
        )

    return {
        "sheet": sheet_name,
        "mode": "scaled_insert_wrap",
        "cleared_errors": cleared,
        "sample_headers_touched": sample_hdrs,
        "intro": intro_reports,
        "photo_start_row": photo_start,
        "photo_anchors": anchors,
        "sample_size": sample_size,
    }


def apply_golden_sheet(ws, sheet_name: str, sample_size: int = 4) -> dict:
    if sheet_name == "ORT":
        from ate.reporting.rules_layout import apply_rules_ort_sheet

        cleared = clear_error_placeholders(ws)
        info = apply_rules_ort_sheet(ws, sample_size=sample_size)
        info["sheet"] = sheet_name
        info["cleared_errors"] = cleared
        return info
    return apply_golden_non_ort(ws, sheet_name, sample_size=sample_size)
