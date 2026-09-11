"""Lab report writers + Test Database photo paste for RS622XK."""
from __future__ import annotations

from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage

from ate.core.paths import artifact_name, screenshot_dir
from ate.reporting.photo_layout import photo_anchor

MYT = timezone(timedelta(hours=8))


def _ctx():
    from ate.core.database import get_context

    return get_context()


def default_lab_report_path() -> Path:
    ctx = _ctx()
    path = ctx.lab_report_path()
    if path.is_file():
        return path
    legacy = Path(r"C:\Users\OoiJianHong\Downloads\RS622XK_Lab_Report_TTSOP.xlsx")
    return legacy if legacy.is_file() else path


def test_db_path(test_folder: str) -> Path:
    path = _ctx().test_folder(test_folder)
    path.mkdir(parents=True, exist_ok=True)
    return path


def ensure_screenshot_dir(test_folder: str = "", dut_index: int | None = None) -> Path:
    if test_folder:
        dest = screenshot_dir(test_folder, dut_index)
        dest.mkdir(parents=True, exist_ok=True)
        return dest
    dest = _ctx().screenshot_dir("ORT")
    dest.mkdir(parents=True, exist_ok=True)
    return dest


def embed_photo(
    ws,
    image_path: str | Path,
    *,
    anchor_cell: str,
    max_width: int = 420,
    max_height: int = 240,
) -> None:
    path = Path(image_path)
    if not path.is_file():
        raise FileNotFoundError(path)
    img = XLImage(str(path))
    if img.width and img.width > max_width:
        scale = max_width / float(img.width)
        img.width = int(img.width * scale)
        img.height = int(img.height * scale)
    if img.height and img.height > max_height:
        scale = max_height / float(img.height)
        img.width = int(img.width * scale)
        img.height = int(img.height * scale)
    ws.add_image(img, anchor_cell)


def update_summary_status(
    workbook_path: Path,
    test_name: str,
    status: str,
) -> None:
    wb = load_workbook(workbook_path)
    if "Summary" not in wb.sheetnames:
        wb.close()
        return
    ws = wb["Summary"]
    needle = test_name.strip().lower()
    for row in ws.iter_rows(min_row=1, max_row=80, max_col=13):
        for cell in row:
            if not cell.value:
                continue
            text = str(cell.value).strip().lower()
            if text == needle or needle in text or text in needle:
                status_cell = ws.cell(cell.row, min(cell.column + 3, 13))
                status_cell.value = status
                wb.save(workbook_path)
                wb.close()
                return
    wb.close()


def _ort_anchor(unit: int, polarity: str, channel: str = "CHA") -> str:
    return photo_anchor("ORT", unit, channel, polarity=polarity)


def place_ort_photos(
    workbook_path: Optional[Path] = None,
    *,
    positive_paths: list[Path],
    negative_paths: list[Path],
    unit_index: int = 1,
    channel: str = "CHA",
) -> Path:
    path = Path(workbook_path or default_lab_report_path())
    if not path.is_file():
        path = default_lab_report_path()
    if not path.is_file():
        raise FileNotFoundError(f"Lab report not found: {path}")

    wb = load_workbook(path)
    if "ORT" not in wb.sheetnames:
        wb.close()
        raise RuntimeError("Sheet 'ORT' missing from lab report")
    ws = wb["ORT"]

    pos = positive_paths[0] if positive_paths else None
    neg = negative_paths[0] if negative_paths else None
    if pos:
        embed_photo(ws, pos, anchor_cell=_ort_anchor(unit_index, "POS", channel))
    if neg:
        embed_photo(ws, neg, anchor_cell=_ort_anchor(unit_index, "NEG", channel))

    try:
        wb.save(path)
    except PermissionError:
        fallback = _ctx().lab_report_path()
        fallback.parent.mkdir(parents=True, exist_ok=True)
        wb.save(fallback)
        path = fallback
    wb.close()
    return path


def _settling_anchor(unit: int, channel: str = "CHA") -> str:
    return photo_anchor("SettlingTime", unit, channel)


def place_settling_photos(
    workbook_path: Optional[Path] = None,
    *,
    photo_path: Path | None,
    unit_index: int = 1,
    channel: str = "CHA",
) -> Path:
    """Paste one Settling photo into the 8-box DUT/channel grid on SettlingTime."""
    if photo_path is None:
        raise ValueError("photo_path required")
    path = Path(workbook_path or default_lab_report_path())
    if not path.is_file():
        path = default_lab_report_path()
    if not path.is_file():
        raise FileNotFoundError(f"Lab report not found: {path}")

    wb = load_workbook(path)
    sheet = "SettlingTime" if "SettlingTime" in wb.sheetnames else "Settling"
    if sheet not in wb.sheetnames:
        wb.close()
        raise RuntimeError("Sheet 'SettlingTime' missing from lab report")
    ws = wb[sheet]
    embed_photo(ws, photo_path, anchor_cell=_settling_anchor(unit_index, channel))
    try:
        wb.save(path)
    except PermissionError:
        fallback = _ctx().lab_report_path()
        fallback.parent.mkdir(parents=True, exist_ok=True)
        wb.save(fallback)
        path = fallback
    wb.close()
    return path


def _place_mapped_photo(
    workbook_path: Optional[Path],
    *,
    test_key: str,
    photo_path: Path,
    unit_index: int,
    channel: str,
    polarity: str | None = None,
) -> Path:
    entry = _ctx().test_entry(test_key)
    sheet_name = str((entry or {}).get("excel_sheet") or test_key)
    path = Path(workbook_path or default_lab_report_path())
    if not path.is_file():
        path = default_lab_report_path()
    if not path.is_file():
        raise FileNotFoundError(f"Lab report not found: {path}")

    wb = load_workbook(path)
    if sheet_name not in wb.sheetnames:
        wb.close()
        raise RuntimeError(f"Sheet {sheet_name!r} missing from lab report")
    ws = wb[sheet_name]
    embed_photo(
        ws,
        photo_path,
        anchor_cell=photo_anchor(test_key, unit_index, channel, polarity),
    )
    try:
        wb.save(path)
    except PermissionError:
        fallback = _ctx().lab_report_path()
        fallback.parent.mkdir(parents=True, exist_ok=True)
        wb.save(fallback)
        path = fallback
    wb.close()
    return path


def place_mapped_photos(
    workbook_path: Optional[Path] = None,
    *,
    test_key: str,
    photo_path: Path | None,
    unit_index: int = 1,
    channel: str = "CHA",
    polarity: str | None = None,
) -> Path:
    """Paste one photo using sheet_map tests.<test_key>.paste.photos."""
    if photo_path is None:
        raise ValueError("photo_path required")
    return _place_mapped_photo(
        workbook_path,
        test_key=test_key,
        photo_path=photo_path,
        unit_index=unit_index,
        channel=channel,
        polarity=polarity,
    )


def place_sssr_photos(
    workbook_path: Optional[Path] = None,
    *,
    photo_path: Path | None,
    unit_index: int = 1,
    channel: str = "CHA",
) -> Path:
    """Paste one SSSR photo using sheet_map SmallSignalStep paste.photos."""
    if photo_path is None:
        raise ValueError("photo_path required")
    return _place_mapped_photo(
        workbook_path,
        test_key="SmallSignalStep",
        photo_path=photo_path,
        unit_index=unit_index,
        channel=channel,
    )


def place_lssr_photos(
    workbook_path: Optional[Path] = None,
    *,
    photo_path: Path | None,
    unit_index: int = 1,
    channel: str = "CHA",
) -> Path:
    """Paste one LSSR photo using sheet_map LargeSignalStep paste.photos."""
    if photo_path is None:
        raise ValueError("photo_path required")
    return _place_mapped_photo(
        workbook_path,
        test_key="LargeSignalStep",
        photo_path=photo_path,
        unit_index=unit_index,
        channel=channel,
    )


def write_ort_result(
    workbook_path: Optional[Path] = None,
    *,
    unit_index: int,
    polarity: str,
    channel: str,
    value_us: float,
) -> Path:
    """Write ORT µs into right-top result grid (R13:U13 / AA13:AD13 / R20:U20 / AA20:AD20)."""
    path = Path(workbook_path or default_lab_report_path())
    wb = load_workbook(path)
    ws = wb["ORT"]
    unit = max(1, min(4, int(unit_index)))
    col_cha = {1: "R", 2: "S", 3: "T", 4: "U"}[unit]
    col_chb = {1: "AA", 2: "AB", 3: "AC", 4: "AD"}[unit]
    row = 13 if polarity.upper().startswith("POS") else 20
    col = col_cha if "A" in channel.upper() else col_chb
    ws[f"{col}{row}"] = float(value_us)
    wb.save(path)
    wb.close()
    return path


def save_named_screenshot(
    test_folder: str,
    *,
    unit: int,
    variant: str,
    source_bytes_path: Path,
    channel: str = "",
) -> Path:
    dest_dir = ensure_screenshot_dir(test_folder, unit)
    ts = datetime.now(MYT).strftime("%Y-%m-%d_%H%M%S")
    tag = f"{variant}_{channel}" if channel else variant
    name = artifact_name(test_folder, unit, tag, timestamp=ts, ext=Path(source_bytes_path).suffix)
    dest = dest_dir / name
    dest.write_bytes(Path(source_bytes_path).read_bytes())
    return dest
