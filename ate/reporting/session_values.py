"""Fill campaign workbook number cells from living report.json.

Uses sheet_map paste.values (measurement id -> cell, DUT list, or CHA/CHB grid).
paste.narrative (probed intro B cells) gets datasheet/conditions/conclusion text.
Skips FILL_ME. Does not run OpAmp golden layout on other families.
#VALUE! on the filled copy is cleared. Excel comments mark paste cells.
"""
from __future__ import annotations

import logging
import re
import shutil
from pathlib import Path
from typing import Any

from ate.reporting.session_paste import _is_real_cell

from openpyxl.cell.cell import MergedCell
from openpyxl.comments import Comment

_LOG = logging.getLogger("ate.session_values")
_CELL_RE = re.compile(r"^[A-Z]{1,3}\d{1,5}$")
_ERROR_TOKENS = {"#VALUE!", "#REF!", "#N/A", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!"}
_ATE_NOTE = "ATE"


def _cell_token(raw: Any) -> str:
    s = str(raw or "").split("#", 1)[0].strip().upper()
    return s if _CELL_RE.match(s) else ""


def _sheet_and_cell(token: str, default_sheet: str) -> tuple[str, str]:
    if "!" in token:
        sh, cell = token.split("!", 1)
        return sh.strip() or default_sheet, cell.strip()
    return default_sheet, token


def _dut_index(step: dict[str, Any]) -> int:
    try:
        return max(1, int(step.get("dut") or 1))
    except (TypeError, ValueError):
        return 1


def _channel_key(step: dict[str, Any]) -> str:
    raw = str(step.get("channel") or "CHA").strip().upper()
    if raw in ("CHB", "B", "2"):
        return "CHB"
    return "CHA"


def _expand_cells(raw: Any, dut: int, channel: str) -> list[str]:
    """One cell for this DUT/channel. List is DUT-indexed; dict is CHA/CHB."""
    if raw is None:
        return []
    if isinstance(raw, dict):
        pick = (
            raw.get(channel)
            or raw.get(channel.lower())
            or raw.get("CHA")
            or raw.get("cha")
        )
        return _expand_cells(pick, dut, channel)
    if isinstance(raw, (list, tuple)):
        i = dut - 1
        if i < 0 or i >= len(raw):
            return []
        item = raw[i]
        if not _is_real_cell(item):
            return []
        tok = _cell_token(item)
        return [tok] if tok else []
    if not _is_real_cell(raw):
        return []
    tok = _cell_token(raw)
    return [tok] if tok else []


def _fold_name(raw: str) -> set[str]:
    """voh_load -> {vohload, voh}; Supply_Current -> {supplycurrent}."""
    n = re.sub(r"[^a-z0-9]+", "", str(raw or "").lower())
    if not n:
        return set()
    out = {n}
    for suf in ("load", "sweep"):
        if n.endswith(suf) and len(n) > len(suf):
            out.add(n[: -len(suf)])
    if n in ("idd", "icc"):
        out.add("supplycurrent")
    return out


def _entry_for_step(tests: dict[str, Any], step: dict[str, Any]) -> tuple[str, dict[str, Any]] | None:
    tid = str(step.get("test_id") or "").strip()
    if not tid:
        return None
    want = [tid, str(step.get("lab_sheet") or ""), str(step.get("fixture_mode") or "")]
    want_fold = set()
    for a in want:
        want_fold |= _fold_name(a)
    for key, entry in tests.items():
        if not isinstance(entry, dict):
            continue
        names = [str(key), str(entry.get("folder") or ""), str(entry.get("excel_sheet") or "")]
        name_fold = set()
        for b in names:
            name_fold |= _fold_name(b)
        if want_fold & name_fold:
            return str(key), entry
        for a in want:
            if not a:
                continue
            for b in names:
                if a.lower() == b.lower() or a.lower().replace("_", "") == b.lower().replace("_", ""):
                    return str(key), entry
    return None


def _golden_dest(path: Path, *, demo: bool, copy_golden: bool) -> Path:
    if demo:
        return path.with_name(path.stem + "_demo.xlsx")
    if copy_golden:
        return path.with_name(path.stem + "_filled.xlsx")
    return path


def _sweep_rows(step: dict[str, Any]) -> list[dict[str, Any]]:
    data = step.get("data") if isinstance(step.get("data"), dict) else {}
    rows = data.get("rows") if isinstance(data, dict) else None
    if not isinstance(rows, list):
        return []
    return [r for r in rows if isinstance(r, dict)]


def _write_sweep_sheet(wb, steps: list[Any]) -> int:
    """Preset Sweep sheet on the filled copy. Sorted VCC then freq then AWG."""
    blocks: list[dict[str, Any]] = []
    for step in steps:
        if not isinstance(step, dict):
            continue
        rows = _sweep_rows(step)
        if not rows:
            continue
        blocks.append(
            {
                "test_id": str(step.get("test_id") or ""),
                "dut": step.get("dut") or 1,
                "channel": str(step.get("channel") or "CHA"),
                "rows": rows,
            }
        )
    if not blocks:
        return 0
    ws = wb["Sweep"] if "Sweep" in wb.sheetnames else wb.create_sheet("Sweep")
    ws.delete_rows(1, ws.max_row)
    headers = ["test_id", "dut", "channel"]
    for b in blocks:
        for r in b["rows"]:
            for k in r:
                if str(k) not in headers:
                    headers.append(str(k))
    for col, h in enumerate(headers, 1):
        ws.cell(1, col).value = h
    n = 0
    rix = 2
    def _sk(row: dict[str, Any]) -> tuple:
        try:
            v = float(row.get("VCC") or row.get("vcc") or 0)
        except (TypeError, ValueError):
            v = 0.0
        try:
            f = float(row.get("freq_hz") or 0)
        except (TypeError, ValueError):
            f = 0.0
        return (v, f, str(row.get("AWG") or ""))

    for b in blocks:
        for row in sorted(b["rows"], key=_sk):
            ws.cell(rix, 1).value = b["test_id"]
            ws.cell(rix, 2).value = b["dut"]
            ws.cell(rix, 3).value = b["channel"]
            for col, h in enumerate(headers, 1):
                if col <= 3:
                    continue
                ws.cell(rix, col).value = row.get(h)
            rix += 1
            n += 1
    return n


def _is_error_or_empty(raw: Any) -> bool:
    if raw is None:
        return True
    s = str(raw).strip()
    return (not s) or s.upper() in _ERROR_TOKENS


def _annotate(cell, text: str, *, allow: bool) -> bool:
    if not allow:
        return False
    existing = getattr(cell, "comment", None)
    if existing is not None and str(getattr(existing, "author", "") or "") not in ("", _ATE_NOTE):
        return False
    cell.comment = Comment(str(text)[:900], _ATE_NOTE)
    return True


def _spec_line(spec: dict[str, Any]) -> str:
    bits = [str(spec.get("id") or "")]
    unit = str(spec.get("unit") or "").strip()
    for label, key in (("min", "min"), ("typ", "typ"), ("max", "max")):
        val = spec.get(key)
        if val is None or val == "":
            continue
        bits.append(f"{label} {val}{(' ' + unit) if unit else ''}")
    cond = str(spec.get("conditions") or spec.get("notes") or "").strip()
    if cond:
        bits.append(cond)
    return " ".join(x for x in bits if x)


def _vs_limit(value: Any, spec: dict[str, Any]) -> str:
    try:
        num = float(value)
    except (TypeError, ValueError):
        return "unspec"
    mn, mx = spec.get("min"), spec.get("max")
    if mn is not None:
        try:
            if num < float(mn):
                return "FAIL"
        except (TypeError, ValueError):
            pass
    if mx is not None:
        try:
            if num > float(mx):
                return "FAIL"
        except (TypeError, ValueError):
            pass
    if mn is None and mx is None:
        return "unspec"
    return "PASS"


def _conditions_text(part_key: str, test_id: str) -> str:
    from ate.core.specs import load_part_yaml, test_info_map

    info = (test_info_map(part_key) or {}).get(test_id) or {}
    part = load_part_yaml(part_key)
    modes = part.get("fixture_modes") if isinstance(part, dict) else {}
    if isinstance(modes, dict):
        for meta in modes.values():
            if not isinstance(meta, dict):
                continue
            tests = meta.get("tests") or []
            if test_id not in tests and test_id.replace("_", "") not in [
                str(t).replace("_", "") for t in tests
            ]:
                continue
            lines = [str(meta.get("label") or "").strip()]
            for item in meta.get("checklist") or []:
                if str(item).strip():
                    lines.append(str(item).strip())
            text = "\n".join(x for x in lines if x)
            if text:
                return text
    desc = str(info.get("description") or "").strip()
    return desc


def _datasheet_text(part_key: str, test_id: str, mids: list[str]) -> str:
    from ate.core.specs import load_part_datasheet, load_part_specs, test_info_map

    info = (test_info_map(part_key) or {}).get(test_id) or {}
    specs = load_part_specs(part_key)
    want = {str(x) for x in mids if x}
    lines: list[str] = []
    desc = str(info.get("description") or "").strip()
    if desc:
        lines.append(desc)
    ds = load_part_datasheet(part_key)
    src = str((ds or {}).get("file") or (ds or {}).get("extract") or "").strip()
    picked = [
        s
        for s in specs
        if isinstance(s, dict)
        and (
            str(s.get("test") or "") == test_id
            or str(s.get("id") or "") in want
        )
    ]
    if not picked:
        picked = [s for s in specs if isinstance(s, dict) and str(s.get("id") or "") in want]
    for spec in picked:
        lines.append(_spec_line(spec))
    if src:
        lines.append(f"source: {Path(src).name}")
    return "\n".join(lines)


def _conclusion_text(
    part_key: str,
    test_id: str,
    measurements: list[dict[str, Any]],
) -> str:
    from ate.core.specs import load_part_specs

    specs = {str(s.get("id")): s for s in load_part_specs(part_key) if isinstance(s, dict) and s.get("id")}
    lines: list[str] = []
    for m in measurements:
        if not isinstance(m, dict):
            continue
        mid = str(m.get("id") or "")
        spec = specs.get(mid)
        if not spec:
            continue
        val = m.get("value")
        unit = str(m.get("unit") or spec.get("unit") or "")
        verdict = _vs_limit(val, spec)
        bits = _spec_line(spec)
        lines.append(f"{mid}={val} {unit} vs {bits}: {verdict}".strip())
    if not lines:
        return ""
    return "\n".join(lines)


def _clear_errors(ws) -> int:
    n = 0
    for row in ws.iter_rows(min_row=1, max_row=min(200, ws.max_row or 1), max_col=40):
        for cell in row:
            try:
                if cell.value is not None and str(cell.value).strip().upper() in _ERROR_TOKENS:
                    cell.value = None
                    n += 1
            except AttributeError:
                pass
    return n


def _write_narrative(
    wb,
    tests: dict[str, Any],
    steps: list[dict[str, Any]],
    *,
    part_key: str,
    annotate: bool,
) -> tuple[int, int]:
    wrote = 0
    notes = 0
    by_test: dict[str, list[dict[str, Any]]] = {}
    for step in steps:
        hit = _entry_for_step(tests, step)
        if not hit:
            continue
        key, _entry = hit
        by_test.setdefault(key, []).append(step)
    for key, entry in tests.items():
        if not isinstance(entry, dict):
            continue
        paste = entry.get("paste") if isinstance(entry.get("paste"), dict) else {}
        narrative = paste.get("narrative") if isinstance(paste.get("narrative"), dict) else {}
        if not narrative:
            continue
        sheet = str(entry.get("excel_sheet") or key)
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        group = by_test.get(key) or []
        test_id = ""
        mids: list[str] = []
        meas: list[dict[str, Any]] = []
        for step in group:
            test_id = test_id or str(step.get("test_id") or key)
            for m in step.get("measurements") or []:
                if isinstance(m, dict):
                    meas.append(m)
                    if m.get("id"):
                        mids.append(str(m.get("id")))
        test_id = test_id or str(key)
        blobs = {
            "conditions": _conditions_text(part_key, test_id),
            "datasheet": _datasheet_text(part_key, test_id, mids),
            "conclusion": _conclusion_text(part_key, test_id, meas),
        }
        for kind, raw_cell in narrative.items():
            cell_tok = _cell_token(raw_cell)
            if not cell_tok:
                continue
            sh, coord = _sheet_and_cell(cell_tok, sheet)
            if sh not in wb.sheetnames:
                continue
            target = wb[sh][coord]
            if isinstance(target, MergedCell):
                continue
            text = blobs.get(str(kind), "")
            if str(kind) == "circuitry":
                if _is_error_or_empty(target.value):
                    target.value = None
                    wrote += 1
                if _annotate(
                    target,
                    "Test Circuitry: paste schematic from senior Product Testing Report "
                    "(OpAmp golden) or ask an AI agent. Do not invent a drawing.",
                    allow=annotate,
                ):
                    notes += 1
                continue
            if text and _is_error_or_empty(target.value):
                target.value = text
                wrote += 1
            hint = {
                "conditions": "paste.narrative conditions -- fixture checklist + test_info",
                "datasheet": "paste.narrative datasheet -- limits yaml min/typ/max (not #VALUE!)",
                "conclusion": "paste.narrative conclusion -- measured vs datasheet PASS/FAIL",
            }.get(str(kind), f"paste.narrative {kind}")
            if _annotate(target, hint, allow=annotate):
                notes += 1
        _clear_errors(ws)
    return wrote, notes


def fill_workbook_from_report(
    *,
    report: dict[str, Any] | None = None,
    workbook_path: Path | None = None,
    ctx=None,
    demo: bool = False,
    copy_golden: bool = False,
) -> dict[str, Any]:
    from ate.core.database import get_context
    from ate.core.datalog import load_report

    c = ctx or get_context()
    doc = report if isinstance(report, dict) else load_report(ctx=c)
    from ate.tests.logic.excel_lock import (
        OrphanWorkbook,
        uses_excel_lock,
        write_path_b_workbook,
    )
    from ate.tests.logic.product_model import load_product_model

    model = load_product_model(str(getattr(c, "part_key", "") or ""))
    if uses_excel_lock(model):
        try:
            return write_path_b_workbook(ctx=c, model=model, report=doc)
        except OrphanWorkbook as exc:
            return {
                "filled": 0,
                "skipped": 0,
                "status": "orphan",
                "error": str(exc),
            }
    sm = c.load_sheet_map() if hasattr(c, "load_sheet_map") else {}
    tests = sm.get("tests") if isinstance(sm, dict) else {}
    if not isinstance(tests, dict):
        return {"filled": 0, "skipped": 0, "status": "no_sheet_map"}
    src = Path(workbook_path) if workbook_path else c.lab_report_path()
    if not src.is_file():
        try:
            from ate.core.provision_operator import create_clean_golden_workbook

            ident = c.identity() if hasattr(c, "identity") else {}
            root = Path(ident.get("root") or src.parent.parent)
            made = create_clean_golden_workbook(
                root,
                part=str(ident.get("part") or ""),
                package=str(ident.get("package") or ""),
                operator=str(ident.get("operator") or ""),
                sample_size=int(ident.get("sample_size") or 4),
            )
            cand = Path(str(made.get("path") or ""))
            if cand.is_file():
                src = cand
        except Exception:
            pass
    if not src.is_file():
        return {"filled": 0, "skipped": 0, "status": "no_workbook", "excel": str(src)}
    path = _golden_dest(src, demo=demo, copy_golden=copy_golden)
    if path.resolve() != src.resolve():
        try:
            shutil.copy2(src, path)
        except Exception as exc:
            return {
                "filled": 0,
                "skipped": 0,
                "status": "error",
                "excel": str(src),
                "error": str(exc),
            }

    writes: list[tuple[str, str, Any]] = []
    skipped = 0
    for step in doc.get("steps") or []:
        if not isinstance(step, dict):
            continue
        hit = _entry_for_step(tests, step)
        if not hit:
            skipped += 1
            continue
        _key, entry = hit
        paste = entry.get("paste") if isinstance(entry.get("paste"), dict) else {}
        sheet = str(entry.get("excel_sheet") or _key)
        values = paste.get("values") if isinstance(paste.get("values"), dict) else {}
        meas = step.get("measurements") if isinstance(step.get("measurements"), list) else []
        if not values or not meas:
            skipped += 1
            continue
        dut = _dut_index(step)
        channel = _channel_key(step)
        result_written = False
        for m in meas:
            if not isinstance(m, dict):
                continue
            mid = str(m.get("id") or "")
            if mid and mid in values:
                for cell in _expand_cells(values[mid], dut, channel):
                    writes.append((*_sheet_and_cell(cell, sheet), m.get("value")))
            if result_written:
                continue
            for cell in _expand_cells(values.get("result"), dut, channel):
                if m.get("result"):
                    writes.append((*_sheet_and_cell(cell, sheet), str(m.get("result")).upper()))
                    result_written = True

    steps = [s for s in (doc.get("steps") or []) if isinstance(s, dict)]
    has_sweep = any(_sweep_rows(s) for s in steps)
    has_narrative = any(
        isinstance(e, dict)
        and isinstance((e.get("paste") or {}).get("narrative"), dict)
        and (e.get("paste") or {}).get("narrative")
        for e in tests.values()
    )
    if not writes and not has_sweep and not has_narrative:
        return {"filled": 0, "skipped": skipped, "status": "nothing_mapped", "excel": str(path)}

    from openpyxl import load_workbook

    wb = None
    saved = path
    status = "ok"
    notes = 0
    narrative_n = 0
    photos_n = 0
    annotate = True
    part_key = str(getattr(c, "part_key", "") or "")
    on_copy = path.resolve() != src.resolve()
    try:
        wb = load_workbook(path)
        n = 0
        for sh, cell, val in writes:
            if sh not in wb.sheetnames or not cell:
                skipped += 1
                continue
            target = wb[sh][cell]
            if isinstance(target, MergedCell):
                skipped += 1
                continue
            target.value = val
            n += 1
            if _annotate(target, f"paste.values {sh}!{cell} from report.json", allow=annotate):
                notes += 1
        if has_sweep and (demo or copy_golden or "Sweep" in wb.sheetnames):
            n += _write_sweep_sheet(wb, steps)
        nar, nar_notes = _write_narrative(
            wb, tests, steps, part_key=part_key, annotate=annotate
        )
        narrative_n = nar
        notes += nar_notes
        n += nar
        try:
            wb.save(path)
        except PermissionError:
            alt = path.with_name(path.stem + "_filled.xlsx")
            wb.save(alt)
            saved = alt
            status = "locked"
        try:
            wb.close()
        except Exception:
            pass
        wb = None
        if on_copy or demo or copy_golden:
            try:
                from ate.reporting.session_paste import paste_session_photos

                pasted = paste_session_photos(workbook_path=saved, ctx=c)
                photos_n = int(pasted.get("pasted") or 0)
            except Exception as exc:
                _LOG.warning("fill photos: %s", exc)
        return {
            "filled": n,
            "skipped": skipped,
            "status": status,
            "excel": str(saved),
            "golden_source": str(src),
            "narrative": narrative_n,
            "annotated": notes,
            "photos": photos_n,
        }
    except PermissionError as exc:
        return {"filled": 0, "skipped": skipped, "status": "locked", "excel": str(path), "error": str(exc)}
    except Exception as exc:
        _LOG.warning("fill workbook: %s", exc)
        return {"filled": 0, "skipped": skipped, "status": "error", "excel": str(path), "error": str(exc)}
    finally:
        if wb is not None:
            try:
                wb.close()
            except Exception:
                pass
